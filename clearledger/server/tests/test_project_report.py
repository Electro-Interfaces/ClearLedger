"""Отчёты по проектам: история строками «было — стало», чек-лист, Excel проекта."""
import io
from datetime import datetime, timezone
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services import project_report as pr


def ev(kind, text=None, changes=None, frm=None, to=None):
    return SimpleNamespace(kind=kind, text=text, changes=changes, from_stage=frm, to_stage=to,
                           created_at=datetime(2026, 10, 4, 9, 0, tzinfo=timezone.utc), source="user")


def test_событие_раскладывается_по_изменениям():
    lines = pr.event_lines(ev("edit", "Правка паспорта", [
        {"label": "Мощность", "oldDisplay": "100", "newDisplay": "150"},
        {"label": "Интеграция: сверки", "old": [], "new": [{"period": "ноябрь"}]},
    ]), {})
    assert [(l["field"], l["old"], l["new"]) for l in lines] == [
        ("Мощность", "100", "150"), ("Интеграция: сверки", "[]", '[{"period": "ноябрь"}]')]
    assert all(l["kind"] == "Правка" for l in lines)


def test_смена_стадии_без_изменений_и_заметка():
    assert pr.event_lines(ev("stage", frm="lead", to="screening"), {"lead": "Лид", "screening": "Скрининг"})[0] \
        == {"kind": "Смена стадии", "text": "", "field": "Стадия", "old": "Лид", "new": "Скрининг"}
    note = pr.event_lines(ev("note", "Звонил собственник"), {})
    assert len(note) == 1 and note[0]["text"] == "Звонил собственник" and note[0]["field"] == ""


def test_длинное_значение_усекается():
    line = pr.event_lines(ev("edit", changes=[{"label": "x", "old": "а" * 900, "new": ""}]), {})[0]
    assert len(line["old"]) == 498 and line["old"].endswith("…")


def test_строка_чек_листа():
    assert pr.gate_line("ЭЗС-1", "Т", "Скрининг", {"key": "2.1", "label": "Мощность", "required": True})["holds"] == "держит переход"
    waived = pr.gate_line("ЭЗС-1", "Т", "Скрининг", {"key": "2.1", "required": True, "waived": True, "waivedBy": "Иванов", "waiveReason": "нет смысла ждать"})
    assert (waived["status"], waived["holds"], waived["by"], waived["note"]) == ("Обязательность снята", "", "Иванов", "нет смысла ждать")
    stale = pr.gate_line("ЭЗС-1", "Т", "Переговоры", {"key": "2.2", "done": False, "needsConfirmation": True})
    assert stale["status"] == "Требует повторного подтверждения"


def _report(integration=None):
    return {
        "generatedAt": "04.10.2026 12:00",
        "summary": {"projectNo": "ЭЗС-2026-1", "title": "Тест", "kind": "integration" if integration else "new_build",
                    "stage": "screening", "stageLabel": "Скрининг", "phaseLabel": "Подбор", "owner": "Иванов",
                    "region": "Брянская", "address": "Брянск", "nextAction": "Запросить ТУ", "nextActionDue": "2026-10-10",
                    "stageSince": "2026-09-01", "createdAt": "01.09.2026", "plannedPowerKwt": 150.0, "plannedEzs": 2,
                    "freePowerKwt": 150, "controlForm": "аренда", "contractStart": "", "commissionedOn": "", "exitKind": ""},
        "gate": {"stageLabel": "Скрининг", "blocking": ["Мощность"], "closed": 5, "total": 11},
        "phases": [], "checklist": [pr.gate_line("ЭЗС-2026-1", "Тест", "Скрининг", {"key": "2.1", "label": "Мощность", "required": True})],
        "work": [{"kind": "Поручение", "key": "№12", "work": "Запросить ТУ", "state": "В работе", "responsible": "Петров",
                  "due": "10.10.2026", "overdue": "", "created": "01.10.2026"}],
        "techConnection": {"status": "Заявка подана", "gridOperator": "Брянскэнерго", "powerKwt": 150.0},
        "equipment": [], "costs": [{"kind": "ТП", "capital": "капвложение", "title": "", "plan": 100.0, "fact": None, "docRef": ""}],
        "budget": {"plan": 100.0, "fact": 0.0, "delta": -100.0}, "documents": [],
        "history": [{"at": "04.10.2026 12:00", "projectNo": "ЭЗС-2026-1", "title": "Тест", "author": "Иванов",
                     "kind": "Правка", "text": "", "field": "Мощность", "old": "100", "new": "150"}],
        **({"integration": integration} if integration else {}),
    }


def test_excel_проекта_стройки():
    wb = load_workbook(io.BytesIO(pr.project_xlsx(_report())))
    assert wb.sheetnames == ["Сводка", "Чек-лист", "Работа", "Присоединение", "Оборудование", "Бюджет", "Документы", "История"]
    assert wb["История"]["E2"].value == "Мощность" and wb["История"]["G2"].value == "150"
    assert wb["Присоединение"]["A2"].value == "Статус"
    assert any(c.value == "Держит переход" for c in wb["Сводка"]["A"])


def test_excel_проекта_интеграции():
    integ = {"partner": {"name": "Партнёр"}, "commercial": {}, "settlement": {"period": "ежемесячно"}, "technical": {}, "work": {},
             "scenarios": [], "stations": [], "tests": [], "reconciliations": [], "listVersions": []}
    wb = load_workbook(io.BytesIO(pr.project_xlsx(_report(integ))))
    assert {"Сценарии", "Станции", "Испытания", "Сверки", "Версии перечней", "Паспорт интеграции"} <= set(wb.sheetnames)
    assert "Присоединение" not in wb.sheetnames
    assert wb["Паспорт интеграции"]["C2"].value == "ежемесячно"

def test_стадия_интеграции_подписывается_по_коду_а_не_снимку():
    line = pr.event_lines(ev("stage", changes=[{"field": "stage", "label": "Стадия", "old": "decision", "new": "contracting",
                                                "oldDisplay": "Решение", "newDisplay": "Оформление земли"}]),
                          {"decision": "Решение о пилоте", "contracting": "Пилотное соглашение"})[0]
    assert (line["old"], line["new"]) == ("Решение о пилоте", "Пилотное соглашение")

def test_текст_смены_стадии_интеграции_переподписывается():
    from app.services.ezs_site_work import relabel_stage_text
    assert relabel_stage_text("Решение → Оформление земли", "decision", "contracting", "integration") == "Решение о пилоте → Пилотное соглашение"
    tail = relabel_stage_text("Проработка → Решение. Не закрыто на гейте: X", "dd", "decision", "integration")
    assert tail == "Техническое согласование → Решение о пилоте. Не закрыто на гейте: X"
    assert relabel_stage_text("Лид → Скрининг", "lead", "screening", "new_build") == "Лид → Скрининг"
    assert relabel_stage_text("Ручной текст", "lead", "screening", "integration") == "Ручной текст"

