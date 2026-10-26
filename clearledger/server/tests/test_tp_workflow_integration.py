"""HTTP/SQL приёмка ТП на отдельной локальной БД tp_test (без общего conftest).
TEST_DATABASE_URL=postgresql+asyncpg://...@127.0.0.1:54339/tp_test
PYTHONPATH=server python -m unittest discover -s server/tests -p test_tp_workflow_integration.py -v
"""
import asyncio
import copy
import os
import unittest
import uuid
from datetime import date, timedelta
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import FastAPI
from sqlalchemy import select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.pool import NullPool

from app.auth import get_current_user
from app.database import Base, get_db
from app import models as m
from app.routers.sites_router import router
from app.services import ezs_tp_workflow as tp

URL = os.environ.get('TEST_DATABASE_URL', '')
RUN = bool(URL and make_url(URL).host in ('localhost', '127.0.0.1') and make_url(URL).database == 'tp_test')

@unittest.skipUnless(RUN, 'Требуется отдельная локальная база tp_test')
class TpIntegration(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_async_engine(URL, poolclass=NullPool)
        cls.sessions = async_sessionmaker(cls.engine, expire_on_commit=False)
        names = ['ApprovalRequest', 'OutboxEvent', 'EventSubscription', 'AuditEvent', 'Company', 'User', 'UserCompany', 'Organization', 'Contract', 'EzsSite', 'EzsTechConnection',
                 'TaskEvent', 'WorkContextResult', 'EzsProject', 'EzsSiteDoc', 'EzsSiteParticipant', 'EzsSiteEquipment', 'EzsSiteCost', 'EzsSiteEvent', 'SourceFile', 'TpOrganizationDocument', 'DocCard', 'DocKind',
                 'DocSignatureEvidence', 'DocRelation', 'DocAcquaint', 'DocVersion', 'DocEvent', 'DocAccessGrant', 'ChatMessage', 'ChatRoom', 'ChatParticipant']
        tables = {getattr(m, n).__table__ for n in names}
        while True:
            parents = {fk.column.table for t in tables for fk in t.foreign_keys}
            if parents <= tables: break
            tables |= parents
        async def create():
            async with cls.engine.begin() as c:
                from app.services.phone_e164 import PHONE_FUNCTION_DDL
                await c.execute(text(PHONE_FUNCTION_DDL))
                await c.run_sync(lambda sync: Base.metadata.create_all(sync, tables=list(tables)))
        asyncio.run(create())

    @classmethod
    def tearDownClass(cls):
        async def drop():
            async with cls.engine.begin() as c:
                await c.execute(text('DROP SCHEMA public CASCADE'))
                await c.execute(text('CREATE SCHEMA public'))
            await cls.engine.dispose()
        asyncio.run(drop())

    async def asyncSetUp(self):
        async with self.engine.begin() as c: await c.execute(text('TRUNCATE companies, users CASCADE'))
        async with self.sessions() as db:
            self.company = m.Company(name='Тест ТП', slug='tp-test', profile_id='energy')
            self.foreign = m.Company(name='Чужая', slug='tp-foreign', profile_id='energy')
            db.add_all([self.company, self.foreign]); await db.flush()
            self.user = m.User(name='Тест энергетик', email='tp@example.invalid', password_hash='!', is_superadmin=True)
            self.other = m.User(name='Чужой', email='other@example.invalid', password_hash='!')
            db.add_all([self.user, self.other]); await db.flush()
            db.add_all([m.UserCompany(user_id=self.user.id, company_id=self.company.id, role='admin'),
                        m.UserCompany(user_id=self.other.id, company_id=self.foreign.id)])
            self.org = m.Organization(name='Тест заявитель', company_id=self.company.id, inn='1', ogrn='2',
                                      legal_address='Тест адрес', director_name='Тест', phone='000')
            self.site = m.EzsSite(company_id=self.company.id, title='Тест ТП', stage='construction', address='Тест адрес', kind='new_build')
            db.add_all([self.org, self.site]); await db.flush()
            self.tc = m.EzsTechConnection(company_id=self.company.id, site_id=self.site.id, method='direct',
                                          power_kwt=100, existing_power_kwt=20, grid_operator='Тест сети', energy_supplier='Тест сбыт')
            db.add(self.tc); await db.commit()
        self.app = FastAPI(); self.app.include_router(router, prefix='/api')
        async def actor(): return self.user
        async def session():
            async with self.sessions() as db: yield db
        self.app.dependency_overrides[get_current_user] = actor
        self.app.dependency_overrides[get_db] = session
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url='http://test')
        self.path = f'/api/sites/{self.site.id}/tech-connection'
        self.q = {'company_id': str(self.company.id)}
        self.w = tp.Workflow(organization_id=self.org.id, personal_data_consent='portal',
                             design_date='2026-11-01', commissioning_date='2026-12-01',
                             responsibilities={'default': tp.Responsibility(executor=self.user.id, controller=self.user.id, provider=self.user.id)}).model_dump(mode='json')

    async def asyncTearDown(self): await self.client.aclose()

    async def save(self, w=None, version=0):
        current = await self.client.get(self.path+'/workflow', params=self.q)
        return await self.client.put(self.path+'/workflow', params=self.q, json={'version': version, 'evidence_token': current.json()['evidenceToken'], 'data': w or self.w})

    async def doc(self, kind, shared=False, expired=False):
        if kind == 'contract': await self.doc('tp_layout')
        async with self.sessions() as db:
            f = m.SourceFile(company_id=self.company.id, file_name='test.pdf', storage_path='/tmp/test-tp.pdf', mime_type='application/pdf')
            db.add(f); await db.flush()
            if shared:
                db.add(m.TpOrganizationDocument(company_id=self.company.id, organization_id=self.org.id, file_id=f.id, kind=kind,
                    valid_until='2000-01-01' if expired else None))
            else:
                db.add(m.EzsSiteDoc(company_id=self.company.id, site_id=self.site.id, kind=kind, file_id=f.id))
            await db.commit()

    async def test_version_and_cross_company_assignment(self):
        r = await self.save(); self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual((await self.save()).status_code, 409)
        bad = copy.deepcopy(self.w); bad['responsibilities']['default']['executor'] = str(self.other.id)
        self.assertEqual((await self.save(bad, 1)).status_code, 400)
        r = await self.client.get(self.path+'/workflow', params={'company_id': str(self.foreign.id)})
        self.assertEqual(r.status_code, 404)

    async def test_document_readiness_shared_expiry_and_notification_idempotency(self):
        await self.doc('contract'); await self.doc('egrul', shared=True, expired=True)
        self.w['reviewed_kinds'] = ['contract', 'tp_layout']
        r = await self.save(); self.assertEqual(r.status_code, 200, r.text); self.assertFalse(r.json()['ready'])
        await self.doc('egrul', shared=True)
        self.w['reviewed_kinds'].append('egrul')
        r = await self.save(version=1); self.assertTrue(r.json()['ready'], r.text)
        delivery = AsyncMock(return_value=uuid.uuid4())
        with patch('app.services.notify.notify_person', delivery):
            async with self.sessions() as db:
                tc = await db.get(m.EzsTechConnection, self.tc.id)
                await tp.deliver(db, tc); await tp.deliver(db, tc); await db.commit()
        self.assertEqual(delivery.await_count, 1)
        # Новая версия файла требует новой проверки содержания.
        await self.doc('contract')
        r = await self.client.get(self.path+'/workflow', params=self.q)
        self.assertFalse(r.json()['ready'])

    async def test_suppressed_notification_is_not_confirmed(self):
        await self.doc('contract'); await self.doc('egrul')
        self.w['reviewed_kinds'] = ['contract', 'egrul', 'tp_layout']; await self.save()
        delivery = AsyncMock(return_value=None)
        with patch('app.services.notify.notify_person', delivery):
            async with self.sessions() as db:
                tc = await db.get(m.EzsTechConnection, self.tc.id)
                await tp.deliver(db, tc); self.assertFalse(tc.workflow['_delivery'])
                await tp.deliver(db, tc)
        self.assertEqual(delivery.await_count, 2)

    async def test_increase_redistribution_and_required_agreements(self):
        self.w.update(scenario='redistribution', representative=True, apartment_building=True)
        r=await self.save(); self.assertEqual(r.status_code,200,r.text)
        kinds={x['kind'] for x in r.json()['kit']}
        self.assertTrue({'tp_existing_act','tp_power_share','poa','tp_building_consent'} <= kinds)

    async def test_open_remarks_and_incomplete_archive_rejected(self):
        self.w['archived']=True
        self.assertEqual((await self.save()).status_code,400)
        self.w['archived']=False
        self.w['remarks']=[dict(id=str(uuid.uuid4()), branch='inspection', text='Нет маркировки', received='2026-10-01', due='2026-10-05')]
        r=await self.save(); self.assertEqual(r.status_code,200,r.text)
        r=await self.client.put(self.path, params=self.q, json={'steps': {'checked':'2026-10-20'}})
        self.assertEqual(r.status_code,400,r.text)
        self.w['remarks'][0]['resolved']='2026-10-20'
        self.assertEqual((await self.save(version=1)).status_code,422) # Нет описания ответа.

    async def test_landlord_keeps_supply_branch_and_does_not_change_tp_status(self):
        r=await self.client.put(self.path,params=self.q,json={'method':'landlord','steps':{'supply_applied':'2026-10-20'}})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual({x['branch'] for x in r.json()['steps']},{'supply'})
        self.assertEqual(r.json()['status'],'draft')

    async def test_print_has_selected_scenario_dates_and_capacity_guard(self):
        self.w['scenario']='increase'; await self.save()
        r=await self.client.get(self.path+'/application',params=self.q)
        self.assertEqual(r.status_code,200,r.text)
        for fragment in ('Увеличение мощности','2026-11-01','2026-12-01','Тест заявитель'): self.assertIn(fragment,r.text)
        await self.client.put(self.path,params=self.q,json={'power_kwt':200})
        r=await self.client.get(self.path+'/application',params=self.q)
        self.assertEqual(r.status_code,400)

    async def test_reading_does_not_initialize_workflow(self):
        for _ in range(2):
            r=await self.client.get(self.path+'/workflow',params=self.q); self.assertEqual(r.status_code,200,r.text)
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            self.assertIsNone(tc.workflow); self.assertEqual(tc.workflow_version,0)

    async def test_existing_connection_stays_manual_until_explicit_start(self):
        r = await self.save()
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(r.json()["data"]["process_enabled"])
        self.assertFalse(r.json()["autoStartNew"])
        self.assertFalse(r.json()["autoStartedNew"])

    async def test_new_connection_starts_after_required_people_are_saved(self):
        async with self.sessions() as db:
            tc = await db.get(m.EzsTechConnection, self.tc.id)
            await db.delete(tc)
            await db.commit()
        created = await self.client.put(self.path, params=self.q, json={"method": "direct"})
        self.assertEqual(created.status_code, 200, created.text)
        state = await self.client.get(self.path + "/workflow", params=self.q)
        self.assertTrue(state.json()["autoStartNew"])
        self.assertFalse(state.json()["data"]["process_enabled"])

        incomplete = copy.deepcopy(self.w)
        incomplete["responsibilities"]["default"]["controller"] = None
        pending = await self.client.put(self.path + "/workflow", params=self.q, json={
            "version": state.json()["version"],
            "evidence_token": state.json()["evidenceToken"],
            "data": incomplete,
        })
        self.assertEqual(pending.status_code, 200, pending.text)
        self.assertTrue(pending.json()["autoStartNew"])
        self.assertFalse(pending.json()["data"]["process_enabled"])

        launched = await self.save(self.w, version=pending.json()["version"])
        self.assertEqual(launched.status_code, 200, launched.text)
        self.assertFalse(launched.json()["autoStartNew"])
        self.assertTrue(launched.json()["autoStartedNew"])
        self.assertTrue(launched.json()["data"]["process_enabled"])

    async def test_new_landlord_connection_does_not_start_a_network_process(self):
        async with self.sessions() as db:
            tc = await db.get(m.EzsTechConnection, self.tc.id)
            await db.delete(tc)
            await db.commit()
        created = await self.client.put(self.path, params=self.q, json={"method": "landlord"})
        self.assertEqual(created.status_code, 200, created.text)
        state = await self.client.get(self.path + "/workflow", params=self.q)
        r = await self.client.put(self.path + "/workflow", params=self.q, json={
            "version": state.json()["version"],
            "evidence_token": state.json()["evidenceToken"],
            "data": self.w,
        })
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["autoStartNew"])
        self.assertFalse(r.json()["data"]["process_enabled"])

    async def test_real_chat_notification_and_archive(self):
        await self.doc('contract'); await self.doc('egrul')
        self.w['reviewed_kinds'] = ['contract', 'egrul', 'tp_layout']
        await self.save()
        with patch('app.services.web_push.push_room_async'):
            async with self.sessions() as db:
                tc = await db.get(m.EzsTechConnection, self.tc.id)
                await tp.deliver(db, tc); await db.commit()
                messages = (await db.execute(select(m.ChatMessage))).scalars().all()
                self.assertEqual(len(messages), 1)
                self.assertIn('/projects?mode=projects', messages[0].content)
                await tp.deliver(db, tc); await db.commit()
                self.assertEqual(len((await db.execute(select(m.ChatMessage))).scalars().all()), 1)
        for kind in ('tu','tp_contract','tp_act','tp_avtu','tp_meter_act'): await self.doc(kind)
        async with self.sessions() as db:
            tc = await db.get(m.EzsTechConnection, self.tc.id)
            tc.done_date='2026-10-26'; await db.commit()
        self.w['archived']=True
        r=await self.save(version=1); self.assertEqual(r.status_code,200,r.text)
        async with self.sessions() as db:
            tc = await db.get(m.EzsTechConnection,self.tc.id)
            self.assertTrue(tc.workflow['_archive']['files'])
        r=await self.client.put(self.path,params=self.q,json={'power_kwt':10})
        self.assertEqual(r.status_code,409)
        r=await self.client.put(self.path,params=self.q,json={'method':'landlord'})
        self.assertEqual(r.status_code,409)

    async def test_targeted_notifications_require_explicit_policy_exception(self):
        await self.doc('contract'); await self.doc('egrul')
        async with self.sessions() as db:
            normal=m.User(name='Тест исполнитель',email='executor@example.invalid',password_hash='!')
            db.add(normal); await db.flush()
            db.add(m.UserCompany(user_id=normal.id,company_id=self.company.id,role='user'))
            company=await db.get(m.Company,self.company.id)
            company.customization={'system_notifications_admin_only':True}
            await db.commit()
        self.w['responsibilities']['default']['executor']=str(normal.id)
        self.w['reviewed_kinds']=['contract','egrul','tp_layout']; await self.save()
        with patch('app.services.web_push.push_room_async'):
            async with self.sessions() as db:
                tc=await db.get(m.EzsTechConnection,self.tc.id)
                await tp.deliver(db,tc); self.assertFalse(tc.workflow['_delivery'])
                company=await db.get(m.Company,self.company.id)
                company.customization=dict(company.customization,tp_notifications_to_assignees=True)
                await db.flush()
                await tp.deliver(db,tc); self.assertTrue(tc.workflow['_delivery'])
                from app.services.notify import notify_person
                self.assertIsNone(await notify_person(db,self.company.id,normal,'Обычное системное сообщение'))
                await db.commit()
                self.assertEqual(len((await db.execute(select(m.ChatMessage))).scalars().all()),1)

    async def test_shared_upload_download_and_non_pdf_rejected(self):
        import tempfile
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {'UPLOAD_DIR': folder}):
            await self.save()
            params=dict(self.q,organization_id=str(self.org.id),kind='egrul')
            r=await self.client.post(self.path+'/shared-documents',params=params,
                files={'file': ('test.pdf',b'%PDF-1.4 test fixture','application/pdf')})
            self.assertEqual(r.status_code,201,r.text)
            doc_id=r.json()['id']
            r=await self.client.get(self.path+f'/shared-documents/{doc_id}/file',params=self.q)
            self.assertEqual(r.status_code,200); self.assertTrue(r.content.startswith(b'%PDF-'))
            r=await self.client.post(self.path+'/shared-documents',params=params,
                files={'file': ('test.pdf',b'not a pdf','application/pdf')})
            self.assertEqual(r.status_code,400)

    async def test_capacity_matrix_boundaries(self):
        from app.services.ezs_tp_method import capacity_profile
        from types import SimpleNamespace
        for power, band in ((150,'до 150 кВт'),(151,'свыше 150 до 670 кВт'),(670,'свыше 150 до 670 кВт'),(671,'свыше 670 кВт')):
            tc=SimpleNamespace(power_kwt=power,existing_power_kwt=0,voltage='35 кВ',method='direct')
            self.assertEqual(capacity_profile(tc)['band'],band)
            self.assertIn('tp_single_line',tp.required_kinds(tc,tp.Workflow()))
        tc.voltage='400 В'
        self.assertNotIn('tp_single_line',tp.required_kinds(tc,tp.Workflow()))
        self.assertIn('tp_automation',tp.required_kinds(tc,tp.Workflow(emergency_automation=True)))

    async def test_prefill_uses_passport_and_keeps_manual_values(self):
        async with self.sessions() as db:
            loc=m.ServiceLocation(id='tp-object',company_id=self.company.id,code='test',name='Тест',power_kwt=350,
                extra_metadata={'connectedPowerKwt':45,'connectionVoltageKv':6,'gridOperator':'Из паспорта'})
            db.add(loc); await db.flush()
            site=await db.get(m.EzsSite,self.site.id);site.location_id=loc.id
            tc=await db.get(m.EzsTechConnection,self.tc.id);tc.existing_power_kwt=None
            await db.commit()
        r=await self.client.post(self.path+'/prefill',params=self.q)
        self.assertEqual(r.status_code,200,r.text)
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            self.assertEqual(tc.existing_power_kwt,45)
            self.assertEqual(tc.grid_operator,'Тест сети')
            self.assertEqual(tc.voltage,'6')
        r=await self.client.post(self.path+'/prefill',params=self.q)
        self.assertEqual(r.json()['applied'],[])

    async def test_handover_is_immutable_idempotent_and_personal(self):
        import tempfile
        from pathlib import Path
        from app.services.ezs_tp_method import create_handover,handover_state
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ,{'UPLOAD_DIR':folder}), patch('app.services.web_push.push_room_async'):
            for kind in ('contract','egrul','tu','tp_contract','tp_act','tp_avtu','tp_meter_act'): await self.doc(kind)
            await self.doc('photo')
            async with self.sessions() as db:
                for f in (await db.execute(select(m.SourceFile))).scalars():
                    p=Path(folder)/str(f.id);p.write_bytes(b'%PDF-1.4 test');f.storage_path=str(p)
                tc=await db.get(m.EzsTechConnection,self.tc.id);tc.done_date='2026-10-26'
                await db.commit()
            self.w['archived']=True
            r=await self.save();self.assertEqual(r.status_code,200,r.text)
            async with self.sessions() as db:
                tc=await db.get(m.EzsTechConnection,self.tc.id)
                await create_handover(db,tc);await create_handover(db,tc);await db.commit()
                state=await handover_state(db,tc)
                self.assertEqual(len(state['recipients']),1);self.assertFalse(state['done'])
                self.assertEqual(len((await db.execute(select(m.DocCard))).scalars().all()),1)
                card=(await db.execute(select(m.DocCard))).scalar_one()
                self.assertNotIn('photo',card.attrs['archive']['files'])
            r=await self.client.post(self.path+'/acquaint',params=self.q)
            self.assertEqual(r.status_code,200,r.text)
            r=await self.client.post(self.path+'/acquaint',params=self.q)
            self.assertTrue(r.json()['repeated'])
            async with self.sessions() as db:
                tc=await db.get(m.EzsTechConnection,self.tc.id)
                self.assertTrue((await handover_state(db,tc))['done'])

    async def test_registered_document_files_are_in_packet_evidence(self):
        await self.doc('tp_contract')
        async with self.sessions() as db:
            local=(await db.execute(select(m.EzsSiteDoc).where(m.EzsSiteDoc.kind=='tp_contract'))).scalar_one()
            fid=local.file_id;await db.delete(local)
            kind=m.DocKind(company_id=self.company.id,code='test_tp_contract',name='Договор ТП',gate_key='tp_contract')
            db.add(kind);await db.flush()
            card=m.DocCard(company_id=self.company.id,kind_id=kind.id,title='Договор',subject_ref=f'site:{self.site.id}',has_files=True)
            db.add(card);await db.flush()
            db.add(m.DocVersion(company_id=self.company.id,doc_id=card.id,file_id=fid,file_name='test.pdf',sha256='0'*64))
            await db.flush()
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            kinds,_,signatures=await tp.evidence(db,tc,tp.data_of(tc))
            self.assertIn('tp_contract',kinds);self.assertEqual(signatures['tp_contract'],[str(fid)])

    async def test_supply_needs_document_not_only_date(self):
        self.w['supply_mode']='direct'; await self.save()
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            tc.steps={'supply_contract':{'date':'2026-10-26'}};await db.commit()
            self.assertFalse((await tp.read(db,tc))['supplyComplete'])
        await self.doc('power_contract')
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            self.assertTrue((await tp.read(db,tc))['supplyComplete'])

    async def test_supplier_prefill_uses_selected_energy_contract(self):
        async with self.sessions() as db:
            cp=m.Counterparty(company_id=self.company.id,name='Сбыт из договора',inn='2');db.add(cp);await db.flush()
            contract=m.Contract(company_id=self.company.id,organization_id=self.org.id,counterparty_id=cp.id,type='energy',number='Э-1',date='2026-10-01')
            db.add(contract);tc=await db.get(m.EzsTechConnection,self.tc.id);tc.energy_supplier=None;await db.commit()
        self.w['supply_contract_id']=str(contract.id);self.w['supply_mode']='existing'
        r=await self.save();self.assertEqual(r.status_code,200,r.text)
        r=await self.client.post(self.path+'/prefill',params=self.q);self.assertEqual(r.status_code,200,r.text)
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id);self.assertEqual(tc.energy_supplier,'Сбыт из договора')
            contract=await db.get(m.Contract,contract.id);contract.type='rent';await db.commit()
        r=await self.save(version=2);self.assertEqual(r.status_code,400)

    async def test_process_sync_preserves_subject_and_retries(self):
        from app.services.ezs_tp_method import reconcile
        self.w['process_enabled']=True;self.w['supply_mode']='direct'
        self.assertEqual((await self.save()).status_code,200)
        calls=AsyncMock(return_value={'processId':str(uuid.uuid4()),'currentStep':{'code':'collect','name':'Сбор'},'state':'running'})
        with patch('app.services.projects_process.sync_subject_process',calls),patch('app.services.notify.notify_person',AsyncMock(return_value=uuid.uuid4())):
            async with self.sessions() as db:
                tc=await db.get(m.EzsTechConnection,self.tc.id)
                await reconcile(db,tc);await db.commit()
                self.assertEqual(set(tc.workflow['_process']['items']),{'tp_connection','tp_supply'})
        self.assertEqual(calls.await_count,2)
        for call in calls.await_args_list:
            self.assertEqual(call.args[3],'tech_connection');self.assertEqual(call.args[4],self.tc.id)

    async def test_stages_reject_acceptance_before_work(self):
        self.w['stages']=[dict(id=str(uuid.uuid4()),title='Монтаж',accepted='2026-10-20')]
        self.assertEqual((await self.save()).status_code,422)



    async def test_projection_explicit_deadline_matches_sql_and_risk_filter(self):
        from app.services.ezs_tp_projection import summary, overdue_sql
        from app.services.ezs_sites import _risk_conditions
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            tc.status='applied';tc.steps={'accepted':{'date':'2030-01-01'}}
            tc.workflow=dict(self.w,responsibilities={'contract':{'due':'2030-01-05'}})
            await db.commit()
            fixed=date(2030,1,10)
            view=summary(tc,fixed)
            self.assertTrue(view['overdue'])
            self.assertEqual(next(r['plannedDate'] for r in view['steps'] if r['key']=='contract'),'2030-01-05')
            for done in (None,'2030-01-04'):
                tc.contract_date=done;await db.flush()
                actual=await db.scalar(text('SELECT '+overdue_sql('t',':tp_today')+' FROM ezs_tech_connections t WHERE id=:id'),{'tp_today':fixed,'id':tc.id})
                self.assertEqual(bool(actual),summary(tc,fixed)['overdue'])
            # Реальный отбор использует текущее число и те же сроки.
            tc.contract_date=None;tc.workflow=dict(self.w,responsibilities={'contract':{'due':'2000-01-01'}})
            await db.flush()
            ids=(await db.execute(select(m.EzsSite.id).where(m.EzsSite.company_id==self.company.id,*_risk_conditions('tp_overdue')))).scalars().all()
            self.assertIn(self.site.id,ids)

    async def test_projection_stages_remarks_supply_and_archived_sql_match(self):
        from app.services.ezs_tp_projection import summary, overdue_sql
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            fixed=date(2030,1,10)
            cases=[dict(stages=[{'id':str(uuid.uuid4()),'title':'Кабель','due':'2030-01-02'}]),
                   dict(remarks=[{'id':str(uuid.uuid4()),'text':'Исправить акт','branch':'inspection','received':'2030-01-01'}]),
                   dict(remarks=[{'id':str(uuid.uuid4()),'text':'Приложить схему','branch':'application','received':'2030-01-01'}]),
                   dict(supply_mode='direct',responsibilities={'supply_applied':{'due':'2030-01-02'}}),
                   dict(supply_mode='landlord',responsibilities={'supply_applied':{'due':'2030-01-02'}}),
                   dict(archived=True,responsibilities={'contract':{'due':'2030-01-02'}})]
            for patch_data in cases:
                tc.workflow=dict(self.w,**patch_data);await db.flush()
                actual=await db.scalar(text('SELECT '+overdue_sql('t',':tp_today')+' FROM ezs_tech_connections t WHERE id=:id'),{'tp_today':fixed,'id':tc.id})
                self.assertEqual(bool(actual),summary(tc,fixed)['overdue'],patch_data)

    async def test_projection_next_step_keeps_manual_and_external_priority(self):
        from app.services.ezs_sites import next_steps
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            tc.workflow=dict(self.w,process_enabled=True,responsibilities={'contract':{'due':'2000-01-01'}},_process={'at':'2030-01-02T12:00:00Z','items':{
                'tp_connection':{'stage':{'code':'grid','name':'Получить договор и ТУ'},'state':'running','performer':{'id':'test','name':'Тест энергетик'}}}})
            site=await db.get(m.EzsSite,self.site.id);await db.flush()
            got=await next_steps(db,self.company.id,self.user,[site])
            self.assertEqual(got[str(site.id)]['source'],'tech_connection')
            from app.services.ezs_tp_projection import next_deadline_sql
            from app.services.ezs_sites import next_step_overdue_sql,_risk_conditions
            planned=await db.scalar(text('SELECT '+next_deadline_sql('t')+' FROM ezs_tech_connections t WHERE id=:id'),{'id':tc.id})
            self.assertEqual(planned.isoformat(),got[str(site.id)]['due'])
            late=await db.scalar(text('SELECT '+next_step_overdue_sql('s')+' FROM ezs_sites s WHERE id=:id'),{'id':site.id})
            self.assertTrue(late)
            missing=(await db.execute(select(m.EzsSite.id).where(m.EzsSite.company_id==self.company.id,*_risk_conditions('no_next')))).scalars().all()
            self.assertNotIn(site.id,missing)
            site.next_action='Проверить бюджет';await db.flush()
            self.assertEqual((await next_steps(db,self.company.id,self.user,[site]))[str(site.id)]['source'],'manual')
            site.workspace_data={'external_wait':{'waiting_for':'Ответ собственника','follow_up':'2030-02-01'}};await db.flush()
            self.assertEqual((await next_steps(db,self.company.id,self.user,[site]))[str(site.id)]['source'],'wait')

    async def test_projection_read_keeps_process_history_and_does_not_write(self):
        from app.services.ezs_tp_projection import read
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            tc.workflow=dict(self.w,process_enabled=True,_process={'at':'2030-01-02T12:00:00Z','items':{
                'tp_connection':{'stage':{'code':'grid','name':'Получить договор и ТУ'},'state':'running','performer':{'id':'test','name':'Тест энергетик'},
                    'startedAt':'2030-01-01T00:00:00Z','history':[{'from_code':'collect','to_code':'grid','at':'2030-01-02T00:00:00Z'}]}}})
            await db.commit();before=copy.deepcopy(tc.workflow);version=tc.workflow_version
            view=await read(db,tc)
            self.assertEqual(view['branches'][0]['history'][0]['to_code'],'grid')
            self.assertEqual(view['branches'][0]['performer']['name'],'Тест энергетик')
            self.assertEqual(tc.workflow,before);self.assertEqual(tc.workflow_version,version)
            self.assertFalse(db.dirty)

    async def test_projection_is_shared_by_roadmap_overview_and_company_scope(self):
        from app.routers.project_workspace_router import router as workspace_router
        self.app.include_router(workspace_router,prefix='/api')
        road=await self.client.get(f'/api/sites/{self.site.id}/roadmap',params=self.q)
        self.assertEqual(road.status_code,200,road.text)
        overview=await self.client.get(f'/api/project-workspace/{self.site.id}',params=self.q)
        self.assertEqual(overview.status_code,200,overview.text)
        self.assertEqual(road.json()['techConnection'],overview.json()['tech_connection'])
        denied=await self.client.get(f'/api/project-workspace/{self.site.id}',params={'company_id':str(self.foreign.id)})
        self.assertEqual(denied.status_code,404)

    async def test_projection_reaches_plan_operator_report_and_export(self):
        import io
        from openpyxl import load_workbook
        from app.services.ezs_lifecycle import list_projects
        from app.services.ezs_project import tech_connections_report,tech_connections_by_operator,export_tech_connections_xlsx,export_portfolio_xlsx
        async with self.sessions() as db:
            tc=await db.get(m.EzsTechConnection,self.tc.id)
            tc.workflow=dict(self.w,process_enabled=True,responsibilities={'contract':{'due':'2000-01-01'}},_process={'items':{
                'tp_connection':{'stage':{'code':'grid','name':'Получить договор и ТУ'},'state':'running','performer':{'id':'test','name':'Тест энергетик'}}}})
            p=m.EzsProject(company_id=self.company.id,site_id=self.site.id,kind='retrofit',stage='construction',title='Тест работа')
            db.add(p);await db.commit()
            plan=await list_projects(db,self.company.id,user=self.user)
            self.assertIn('Получить договор',plan[0]['nextAction'])
            self.assertEqual(plan[0]['nextActionDue'],'2000-01-01')
            report=await tech_connections_report(db,self.company.id)
            operators=await tech_connections_by_operator(db,self.company.id)
            self.assertEqual(report['overdue'],1)
            self.assertEqual(sum(x['overdue'] for x in operators['items']),1)
            data=await export_tech_connections_xlsx(db,self.company.id)
            book=load_workbook(io.BytesIO(data),data_only=True)
            self.assertEqual(book.sheetnames,['Присоединение','Шаги ТП и сбыта','Просрочки ТП'])
            self.assertEqual(book['Присоединение'].max_column,21)
            portfolio=load_workbook(io.BytesIO(await export_portfolio_xlsx(db,self.company.id,user=self.user)),data_only=True)
            found=any('Получить договор' in str(cell.value) for sheet in portfolio for row in sheet for cell in row)
            self.assertTrue(found)
