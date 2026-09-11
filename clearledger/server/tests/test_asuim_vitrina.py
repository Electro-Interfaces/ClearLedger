"""Разбор выгрузок витрины АСУиМ ЭЗС — опознание представления и мапперы.

БД не нужна: проверяются чистые функции (`read_asuim_xlsx`, `map_stations`,
`map_payments`). Формат файла воспроизведён по пробной выгрузке 06.08.2026:
данные лежат на ВТОРОМ листе, первый ODBC оставляет пустым.
"""
import io

import openpyxl

from app.services.asuim_normalize import (
    _conn_type,
    _read_sheet,
    detect_sheets,
    _is_phone_stub,
    _keep_filled,
    _org_phone_stub,
    _phone,
    map_payments,
    map_stations,
    read_asuim_xlsx,
)


def _xlsx(headers: list[str], rows: list[list]) -> bytes:
    """Файл в формате ODBC-выгрузки: пустой «Лист1» + данные на «Лист2»."""
    wb = openpyxl.Workbook()
    wb.active.title = "Лист1"
    ws = wb.create_sheet("Лист2")
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STATION_HEADERS = [
    "id_станции", "название", "номер", "серийный_номер", "регион", "город", "улица",
    "номер_дома", "адрес", "широта", "долгота", "статус", "этап", "id_владельца",
    "название_владельца", "бренд", "модель", "протокол_ocpp", "мощность_квт",
    "количество_коннекторов", "средняя_оценка", "процент_успеха", "id_группы",
]


def test_stations_view_detected_on_second_sheet():
    content = _xlsx(STATION_HEADERS, [[
        "000073", "Новая Рига", "643", "000073", "Московская область", "д.Покровское",
        "Центральная", "33", "Московская область, д.Покровское, Центральная, 33",
        55.810532, 37.023645, "Активная", "Active", 3, "РусГидро", "Нартис", None,
        0, None, 2, 0, 0, None,
    ]])
    view, rows = read_asuim_xlsx(content)
    assert view == "stations"
    assert len(rows) == 1


def test_stations_mapper_keeps_empty_fields_empty():
    """Пустая мощность и нулевой рейтинг не должны стереть накопленный паспорт."""
    content = _xlsx(STATION_HEADERS, [[
        "000073", "Новая Рига", "643", "000073", "Московская область", "д.Покровское",
        "Центральная", "33", "адрес", 55.81, 37.02, "Активная", "Active", 3,
        "РусГидро", "Нартис", None, 0, None, 2, 0, 0, None,
    ]])
    _, rows = read_asuim_xlsx(content)
    row = map_stations(rows)[0]
    assert row["partial"] is True             # режим «дозаполнить»
    assert row["power_kwt"] is None           # пусто в выгрузке
    assert row["rating"] is None              # ноль здесь = «не измеряли»
    assert row["success_pct"] is None
    assert row["ocpp_protocol"] == "1.6"      # код 0 → версия 1.6
    assert row["ext_id"] == "000073"
    assert row["connectors_count"] == 2
    assert row["extra"]["asuimOwnerId"] == 3
    assert row["is_test"] is False


def test_stations_mapper_marks_test_stations():
    content = _xlsx(STATION_HEADERS, [[
        "00000000", "ТестТест", "Тест", "0000", "Большой", "Маленький", "Средняя",
        None, "адрес", 44.4, 55.5, "Отключена", "Active", 3, "РусГидро", "ПСС",
        None, 1, None, 1, 0, 0, None,
    ]])
    _, rows = read_asuim_xlsx(content)
    row = map_stations(rows)[0]
    assert row["is_test"] is True
    assert row["ocpp_protocol"] == "2.0"


PAYMENT_HEADERS = [
    "id_платежа", "id_транзакции_банка", "id_сессии", "дата", "сумма_руб.",
    "сумма_холда_руб.", "сумма_возврата_руб", "оплата_картой", "id_типа_операции",
    "тип_операции", "статус", "url_фискального_чека", "id_пользователя",
    "телефон_пользователя",
]


def test_payments_mapper_splits_hold_and_refund():
    """Холд − возврат = списание: складывать холды в выручку нельзя."""
    content = _xlsx(PAYMENT_HEADERS, [[
        139543, "019fc8f9-76f4", 159070, "2026-08-03 19:04:50", 108.68, 192.28, 83.6,
        1, 3, "Оплата картой за зарядку", 2,
        "https://lk.platformaofd.ru/cheque", 14765, "+79941078040",
    ]])
    view, rows = read_asuim_xlsx(content)
    assert view == "payments"
    p = map_payments(rows)[0]
    assert p["payment_ext_id"] == "139543"
    assert p["session_ext_id"] == "159070"
    assert abs(p["hold_amount"] - p["refund_amount"] - p["amount"]) < 0.01
    assert p["paid_at"].year == 2026
    assert p["by_card"] is True
    assert p["receipt_url"].startswith("https://")


USER_HEADERS = [
    "id_пользователя", "логин", "фамилия", "имя", "отчество", "телефон", "email",
    "баланс_руб", "активен", "аватар", "id_организации", "дата_регистрации",
    "месяц_регистрации", "год_регистрации",
]


def test_users_view_detected_and_phone_normalized():
    """Телефон в справочнике клиентов отформатирован, а в сессиях и платежах —
    сплошной строкой. Без приведения справочник ни с чем не соединится."""
    content = _xlsx(USER_HEADERS, [[
        2, "mkiselev", "Лукин", "Вит", "Александрович", "+7(986) 902 19-72",
        "bender2@mail.ru", 997.24, "true", None, 0, "2024-08-07 13:25:15", 8, 2024,
    ]])
    view, rows = read_asuim_xlsx(content)
    assert view == "users"
    assert _phone(rows[0]["телефон"]) == "+79869021972"
    assert _phone("8 (986) 902-19-72") == "+79869021972"


def test_rfid_view_detected():
    content = _xlsx(
        ["id_карты", "uid", "номер", "статус_код", "статус", "id_пользователя",
         "телефон_пользователя"],
        [[1, "26AA5969", "1700", 2, "Активная", 3, "+79087844272"]])
    view, rows = read_asuim_xlsx(content)
    assert view == "rfid"
    assert rows[0]["uid"] == "26AA5969"


def test_unknown_file_is_not_claimed():
    """Выгрузка админпанели не должна опознаваться как витрина — иначе сломается
    прежний индексный разбор."""
    view, _ = read_asuim_xlsx(_xlsx(["Номер", "Название", "OCPP ID"], [["1", "ЭЗС", "x"]]))
    assert view is None


def test_repeat_load_does_not_erase_known_values():
    """Повторная выгрузка не стирает то, что уже загружено.

    Витрина отдаёт данные частями: пустая колонка означает «в этой выгрузке нет»,
    а не «значение удалили». Безусловная перезапись обнуляла телефон и баланс
    клиента на каждой второй загрузке."""
    class Row:
        pass
    o = Row()
    o.phone, o.balance, o.is_active = "+79990000000", 100.0, True
    _keep_filled(o, {"phone": None, "balance": 250.0, "is_active": None})
    assert o.phone == "+79990000000"   # пустое не затирает
    assert o.balance == 250.0          # непустое обновляет
    assert o.is_active is True


def test_new_organizations_do_not_collide_on_empty_phone():
    """Две новых организации в одном файле не роняют прогон.

    `corporate_clients.phone` уникален по компании и NOT NULL, а витрина телефон
    организации не отдаёт. Пустой строкой заводилась только первая, вторая падала
    по уникальному индексу и откатывала весь файл."""
    assert _org_phone_stub("11") != _org_phone_stub("12")
    assert len(_org_phone_stub("x" * 40)) <= 20
    assert _is_phone_stub(_org_phone_stub("11"))
    assert not _is_phone_stub("+79990000000")


def test_payment_amounts_stay_empty_when_column_is_empty():
    """Пустая сумма остаётся пустой в маппере: ноль на обновлении затёр бы платёж.
    Ноль подставляется только при создании строки (поля NOT NULL)."""
    content = _xlsx(PAYMENT_HEADERS, [[
        139544, "019fc8f9-77aa", None, "2026-08-03 19:04:50", None, None, None,
        1, 3, "Оплата картой за зарядку", 0, None, 14765, "8 (994) 107-80-40",
    ]])
    _, rows = read_asuim_xlsx(content)
    p = map_payments(rows)[0]
    assert p["amount"] is None and p["hold_amount"] is None and p["refund_amount"] is None
    # Телефон плательщика приводится к тому же виду, что в справочнике клиентов,
    # иначе связка «платёж → клиент» по телефону не собирается.
    assert p["user_phone"] == "+79941078040"


def test_connector_type_canon_matches_sessions():
    """Паспорт станции и прайс называют разъём так же, как сессия.

    Слева — написания витрины, справа — то, что лежит в
    `charge_sessions.connector_type`. Разойдутся — разрез по типу разъёма и
    сравнение факта с прайсом молча посчитают один разъём двумя."""
    assert _conn_type("CCS2") == "CCS Combo 2"
    assert _conn_type("GBT/DC") == "GB/T DC"
    assert _conn_type("GBT/AC") == "GB/T AC"
    assert _conn_type("Type 1") == "Type 1"
    assert _conn_type("Chademo") == "CHAdeMO"
    assert _conn_type("TYPE2") == "Type 2"
    assert _conn_type("Переходник T2-T1") == "Type 1"
    # Неоднозначное и пустое не выдумываем.
    assert _conn_type("GB/T") == "GB/T"
    assert _conn_type("") is None


def test_batch_orders_files_by_dependency():
    """Пакет грузится в порядке связей, а не в порядке выбора файлов.

    Станции резолвят всё остальное, телефон организации проставляется из её
    клиента, платёж ссылается на сессию — если порядок перепутать, часть строк
    осядет «сиротами» до следующего прогона."""
    from app.services.asuim_normalize import VIEW_ORDER
    pos = {v: i for i, v in enumerate(VIEW_ORDER)}
    assert pos["stations"] < pos["connectors"]
    assert pos["organizations"] < pos["users"]
    assert pos["sessions"] < pos["payments"]
    assert "admins" not in VIEW_ORDER          # не грузим никогда


# ---------------------------------------------------------------------------
# Книга витрины: 14 листов одним файлом (как приходит выгрузка с 28.08.2026)
# ---------------------------------------------------------------------------
BOOK_SESSION_HEADERS = ["id_сессии", "дата_начала", "энергия_квтч", "сумма_руб"]
BOOK_PAYMENT_HEADERS = ["id_платежа", "id_сессии", "сумма_холда_руб.", "дата", "сумма_руб."]


def _book() -> bytes:
    """Книга ODBC: пустой первый лист, дальше представления вперемешку."""
    wb = openpyxl.Workbook()
    wb.active.title = "Лист1"
    for title, headers, rows in (
        ("Лист2", BOOK_SESSION_HEADERS, [["190681", "2026-09-08 10:00:00", 35.6, 400],
                                    ["198101", "2026-09-10 16:39:00", 13.4, 0]]),
        ("Лист3", BOOK_PAYMENT_HEADERS, [["p1", "190681", 500, "2026-09-08 10:20:00", 400],
                                    ["p2", "198101", 0, "2026-09-10 17:00:00", 0]]),
        ("Лист4", STATION_HEADERS, [["000073", "Новая Рига", "643", "000073", "МО", "д.Покровское",
                                     "Центральная", "33", "адрес", 55.81, 37.02, "Активная",
                                     "Active", 3, "РусГидро", "Нартис", None, 0, None, 2, 0, 0, None]]),
    ):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_book_all_sheets_detected():
    """Из книги должны доехать ВСЕ представления, а не первое опознанное.

    `read_asuim_xlsx` возвращает первый лист — на книге это только сессии,
    без платежей и справочников."""
    content = _book()
    view, _ = read_asuim_xlsx(content)
    assert view == "sessions"

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheets = detect_sheets(wb)
        assert sheets == {"sessions": "Лист2", "payments": "Лист3", "stations": "Лист4"}
    finally:
        wb.close()


def test_book_sheet_read_keeps_only_new_days():
    """Сессии и платежи берутся с новых дней: книга отдаёт всю историю каждый раз."""
    wb = openpyxl.load_workbook(io.BytesIO(_book()), read_only=True, data_only=True)
    try:
        assert len(_read_sheet(wb["Лист2"])) == 2
        rows = _read_sheet(wb["Лист2"], "дата_начала", "2026-09-10")
        assert [r["id_сессии"] for r in rows] == ["198101"]
        rows = _read_sheet(wb["Лист3"], "дата", "2026-09-09")
        assert [r["id_платежа"] for r in rows] == ["p2"]
        # Справочник отсечке не подлежит: даты у него нет, берётся целиком.
        assert len(_read_sheet(wb["Лист4"], None, "2026-09-10")) == 1
    finally:
        wb.close()


def test_stations_mapper_marks_customer_test_stands():
    """Стенд «(Тест)» с боевым номером не должен заводиться в реестр объектов."""
    rows = [
        {"id_станции": "756", "название": 'БЦ "Гидропроект" (Тест)', "номер": "756"},
        {"id_станции": "643", "название": "Новая Рига", "номер": "643"},
        {"id_станции": "900", "название": "Тестовая площадка сети", "номер": "900"},
    ]
    marked = {r["ext_id"]: r["is_test"] for r in map_stations(rows)}
    assert marked == {"756": True, "643": False, "900": False}
