"""ROOT-only offline log/XLSX characterization. No application requests or SQL."""
import io
import logging
import os
from pathlib import Path

import pytest
from gunicorn.glogging import SafeAtoms
from openpyxl import Workbook, load_workbook

from app.api_access_logging import OAuthAccessFilter, install_oauth_access_filters, safe_target
from app.services.export_files import XLSX_MIME, xlsx_response

SYNTHETIC = 'SYNTHETIC_CAPABILITY_20261031'
MARKER = '[redacted-capability-target]'
# Fixed source routes/expected outputs, not a copy of the production matching rules.
ROUTES = [
 ('/api/doc-share/{token}', '/api/doc-share/[redacted]'),
 ('/api/doc-share/verify/{token}', '/api/doc-share/verify/[redacted]'),
 ('/api/doc-share/{token}/file/11111111-1111-4111-8111-111111111111', '/api/doc-share/[redacted]/file/11111111-1111-4111-8111-111111111111'),
 ('/api/doc-share/{token}/ack', '/api/doc-share/[redacted]/ack'),
 ('/api/doc-share/{token}/decide', '/api/doc-share/[redacted]/decide'),
 ('/api/invite/{token}', '/api/invite/[redacted]'),
 ('/api/invite/{token}/respond', '/api/invite/[redacted]/respond'),
 ('/api/invite/{token}/propose', '/api/invite/[redacted]/propose'),
 ('/api/invite/{token}/poll', '/api/invite/[redacted]/poll'),
 ('/api/invite/{token}/poll/vote', '/api/invite/[redacted]/poll/vote'),
 ('/api/invite/{token}/ics', '/api/invite/[redacted]/ics'),
 ('/api/work/calendar/feed/{token}.ics', '/api/work/calendar/feed/[redacted].ics'),
]
ROUTE_IDS=['doc-card','doc-verify','doc-file','doc-ack','doc-decide','invite-card','invite-respond',
           'invite-propose','invite-poll','invite-vote','invite-ics','calendar-feed']


@pytest.mark.parametrize('template,expected', ROUTES, ids=ROUTE_IDS)
@pytest.mark.parametrize('token', [SYNTHETIC, '%53YNTHETIC_CAPABILITY_20261031',
    '"SYNTHETIC_CAPABILITY_20261031"', 'СИНТЕТИЧЕСКИЙ_20261031'],
    ids=['plain','percent-token','quoted-token','unicode-token'])
def test_all_source_capability_shapes_fixed_outputs(template, expected, token):
    target=template.replace('{token}',token)
    assert safe_target(target)==expected
    assert safe_target(expected)==expected
    assert SYNTHETIC not in safe_target(target)


@pytest.mark.parametrize('target,expected', [
 ('https://example.invalid/api/doc-share/SYNTHETIC_CAPABILITY_20261031/ack?echo=SYNTHETIC_CAPABILITY_20261031#SYNTHETIC_CAPABILITY_20261031', 'https://example.invalid/api/doc-share/[redacted]/ack'),
 ('http://example.invalid:8080/api/invite/SYNTHETIC_CAPABILITY_20261031/ics?x=1#fragment', 'http://example.invalid:8080/api/invite/[redacted]/ics'),
 ('HTTPS://example.invalid/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031.ics#fragment', 'HTTPS://example.invalid/api/work/calendar/feed/[redacted].ics'),
 ('/api/invite/SYNTHETIC_CAPABILITY_20261031?echo=SYNTHETIC_CAPABILITY_20261031#SYNTHETIC_CAPABILITY_20261031', '/api/invite/[redacted]'),
 ('/api/doc-share/SYNTHETIC_CAPABILITY_20261031/file/11111111-1111-4111-8111-111111111111#fragment', '/api/doc-share/[redacted]/file/11111111-1111-4111-8111-111111111111'),
 ('/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031.ics.ics', '/api/work/calendar/feed/[redacted].ics'),
 ('/api/doc-share/SYNTHETIC%2F_CAPABILITY_20261031/ack', '/api/doc-share/[redacted]/ack'),
 ('/api/invite/SYNTHETIC%3F%23_CAPABILITY_20261031/poll/vote', '/api/invite/[redacted]/poll/vote'),
], ids=['https-referrer','http-port','uppercase-scheme','duplicate-query-fragment','version-suffix',
        'token-dot-ics','encoded-slash-token','encoded-query-fragment-token'])
def test_origin_and_absolute_targets_keep_only_safe_coordinates(target, expected):
    assert safe_target(target)==expected
    assert SYNTHETIC not in safe_target(target)


@pytest.mark.parametrize('target', [
 '/api/doc-share/', '/api/doc-share/verify/ack','/api/doc-share/verify/decide',
 '/api/doc-share/%76erify/SYNTHETIC_CAPABILITY_20261031',
 '/api/doc-share/SYNTHETIC_CAPABILITY_20261031/unknown',
 '/api/invite/SYNTHETIC_CAPABILITY_20261031/poll/extra',
 '/api/invite//ics', '/api/invite/SYNTHETIC_CAPABILITY_20261031\n/ics',
 '/api/work/calendar/feed/.ics','/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031%2Eics',
 '/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031.ics/extra',
 '/api/%64oc-share/SYNTHETIC_CAPABILITY_20261031',
 '%2Fapi%2Finvite%2FSYNTHETIC_CAPABILITY_20261031',
 '/api/work/calendar/%66eed/SYNTHETIC_CAPABILITY_20261031.ics',
 '"/api/invite/SYNTHETIC_CAPABILITY_20261031"',
 ' https://example.invalid/api/invite/SYNTHETIC_CAPABILITY_20261031',
 '//example.invalid/api/invite/SYNTHETIC_CAPABILITY_20261031',
 'ftp://example.invalid/api/invite/SYNTHETIC_CAPABILITY_20261031',
 'https://SYNTHETIC_CAPABILITY_20261031@example.invalid/api/invite/SYNTHETIC_CAPABILITY_20261031',
], ids=['empty-doc','static-collision-ack','static-collision-decide','encoded-static',
        'unknown-doc-tail','unknown-invite-tail','empty-invite','control', 'empty-feed',
        'encoded-extension','extra-feed-tail','encoded-doc-prefix','encoded-root',
        'encoded-feed-prefix','quoted-target','leading-space','network-path','other-scheme','userinfo'])
def test_ambiguous_known_families_fail_closed_without_path_decoding(target):
    assert safe_target(target)==MARKER
    assert SYNTHETIC not in safe_target(target)


@pytest.mark.parametrize('target', [
 '/api/doc-sharing/SYNTHETIC_CAPABILITY_20261031?x=1',
 '/api/invite-extra/SYNTHETIC_CAPABILITY_20261031',
 '/api/work/calendar/feeding/SYNTHETIC_CAPABILITY_20261031.ics',
 '/api/work/calendar/feed', '/api/work/calendar/feed/rotate?view=personal',
 '/api/doc-share', '/api/invite', '/api/files/11111111-1111-4111-8111-111111111111',
 '/ordinary?next=/api/invite/SYNTHETIC_CAPABILITY_20261031',
 'https://example.invalid/ordinary#view', '/API/invite/SYNTHETIC_CAPABILITY_20261031',
 '/api/./invite/SYNTHETIC_CAPABILITY_20261031',
], ids=['doc-near-miss','invite-near-miss','feed-near-miss','feed-bootstrap','feed-rotate',
        'doc-base','invite-base','ordinary-id','nested-query-outside-scope','ordinary-referrer','case-distinct','dot-prefix-outside-scope'])
def test_unmatched_targets_and_static_feed_operations_are_unchanged(target):
    # ROOT msg_7717a23e12f8: intentional logging-only alias privacy expansion.
    if target == '/api/./invite/SYNTHETIC_CAPABILITY_20261031':
        assert safe_target(target)==MARKER
    else:
        assert safe_target(target)==target


@pytest.mark.parametrize('target,expected', [
 ('/api/work/calendar/feed/rotate?token=SYNTHETIC_CAPABILITY_20261031', '/api/work/calendar/feed/rotate'),
 ('/ordinary?token=SYNTHETIC_CAPABILITY_20261031', '/ordinary'),
 ('/api/chat/ws?t%6fken=SYNTHETIC_CAPABILITY_20261031', '/api/chat/ws'),
 ('https://example.invalid/cb?code=SYNTHETIC_CAPABILITY_20261031', 'https://example.invalid/cb'),
], ids=['rotate-auth-query','old-auth-query','old-ws-query','old-absolute-query'])
def test_existing_query_authority_policy_remains_effective(target, expected):
    assert safe_target(target)==expected


@pytest.mark.parametrize('target',[None,17,b'/api/invite/SYNTHETIC_CAPABILITY_20261031',{'path':'synthetic'}],
                         ids=['none','integer','bytes','mapping'])
def test_nonstring_target_identity(target):
    assert safe_target(target) is target


def emitted(record, formatter):
    stream=io.StringIO()
    handler=logging.StreamHandler(stream);handler.setFormatter(logging.Formatter(formatter))
    logger=logging.Logger(record.name);logger.propagate=False
    logger.addHandler(handler);logger.addFilter(OAuthAccessFilter())
    try:logger.handle(record)
    finally:logger.removeHandler(handler);handler.close()
    return stream.getvalue().rstrip('\n')


@pytest.mark.parametrize('target,expected', [
 ('/api/doc-share/SYNTHETIC_CAPABILITY_20261031/ack', '/api/doc-share/[redacted]/ack'),
 ('/api/invite/SYNTHETIC_CAPABILITY_20261031/ics', '/api/invite/[redacted]/ics'),
 ('/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031.ics', '/api/work/calendar/feed/[redacted].ics'),
],ids=['doc','invite','feed'])
def test_actual_uvicorn_http_log_emission(target, expected):
    record=logging.LogRecord('uvicorn.access',logging.INFO,__file__,1,'%s - "%s %s HTTP/%s" %d',
                             ('client','GET',target+'?echo='+SYNTHETIC,'1.1',200),None)
    assert emitted(record,'%(message)s')==f'client - "GET {expected} HTTP/1.1" 200'


@pytest.mark.parametrize('message,status', [
 ('%s - "WebSocket %s" [accepted]',None),('%s - "WebSocket %s" 403',None),
 ('%s - "WebSocket %s" %d',401)],ids=['accepted','refused','http-denial'])
def test_known_ws_logger_formats_emit_only_safe_target(message,status):
    args=('client','/api/invite/'+SYNTHETIC+'/ics')+((status,) if status is not None else ())
    record=logging.LogRecord('uvicorn.error',logging.INFO,__file__,1,message,args,None)
    expected=message % (('client','/api/invite/[redacted]/ics')+((status,) if status is not None else ()))
    assert emitted(record,'%(message)s')==expected
    assert record.args[0]=='client' and record.args[2:]==args[2:]


@pytest.mark.parametrize('safeatoms',[False,True],ids=['plain-mapping','SafeAtoms'])
@pytest.mark.parametrize('target,expected', [
 ('/api/doc-share/SYNTHETIC_CAPABILITY_20261031/file/11111111-1111-4111-8111-111111111111','/api/doc-share/[redacted]/file/11111111-1111-4111-8111-111111111111'),
 ('/api/invite/"SYNTHETIC_CAPABILITY_20261031"/ics','/api/invite/[redacted]/ics'),
 ('/api/work/calendar/feed/SYNTHETIC_CAPABILITY_20261031.ics','/api/work/calendar/feed/[redacted].ics'),
],ids=['doc-file','quoted-invite','feed'])
def test_gunicorn_U_r_q_f_and_custom_referer_all_propagate_safe_target(safeatoms,target,expected):
    values={'U':target,'q':'echo='+SYNTHETIC,'r':'GET '+target+'?echo='+SYNTHETIC+' HTTP/1.1',
        'm':'GET','H':'HTTP/1.1','f':'https://example.invalid'+target+'#'+SYNTHETIC,
        '{referer}i':'https://example.invalid'+target+'?echo='+SYNTHETIC,
        '{authorization}i':SYNTHETIC,'{cookie}i':SYNTHETIC,'{set-cookie}o':SYNTHETIC,
        '{custom}i':'quoted"value'}
    atoms=SafeAtoms(values) if safeatoms else values
    before=dict(atoms)
    record=logging.LogRecord('gunicorn.access',logging.INFO,__file__,1,
        '%(U)s|%(r)s|%(q)s|%(f)s|%({referer}i)s|%({authorization}i)s|%({cookie}i)s|%({set-cookie}o)s', (atoms,),None)
    message=emitted(record,'%(message)s')
    assert message=='|'.join([expected,'GET '+expected+' HTTP/1.1','',
        'https://example.invalid'+expected,'https://example.invalid'+expected,
        '[redacted]','[redacted]','[redacted]'])
    assert SYNTHETIC not in message and dict(atoms)==before
    assert type(record.args) is type(atoms) and record.args is not atoms
    if safeatoms:
        assert record.args['{ABSENT}i']=='-' and record.args['absent']=='-'
        assert record.args['{Custom}i']=='quoted\\"value'


def test_gunicorn_missing_U_keeps_SafeAtoms_missing_semantics():
    atoms=SafeAtoms({'m':'GET','q':'token='+SYNTHETIC,'H':'HTTP/1.1'})
    record=logging.LogRecord('gunicorn.access',logging.INFO,__file__,1,'%(U)s|%(r)s|%(q)s|%(absent)s',(atoms,),None)
    assert emitted(record,'%(message)s')=='-|GET - HTTP/1.1||-'
    assert 'U' not in record.args and atoms['q']=='token='+SYNTHETIC


@pytest.mark.parametrize('kind',['custom-error','wrong-arity','preformatted','unknown-logger'],
                         ids=['custom-error','wrong-arity','preformatted','unknown-logger'])
def test_unsupported_error_formats_keep_existing_boundary(kind):
    target='/api/invite/'+SYNTHETIC
    if kind=='custom-error':name,msg,args='uvicorn.error','Custom %s %s',('client',target)
    elif kind=='wrong-arity':name,msg,args='uvicorn.error','%s - "WebSocket %s" [accepted]',('client',)
    elif kind=='preformatted':name,msg,args='uvicorn.error','preformatted '+target,()
    else:name,msg,args='custom.error','%s - "WebSocket %s" [accepted]',('client',target)
    record=logging.LogRecord(name,logging.INFO,__file__,1,msg,args,None);old_args=record.args
    assert OAuthAccessFilter().filter(record)
    assert record.args is old_args


def test_filter_installation_restores_global_logger_state():
    names=('uvicorn.access','gunicorn.access','uvicorn.error')
    saved={n:list(logging.getLogger(n).filters) for n in names}
    try:
        for n in names:logging.getLogger(n).filters=[]
        install_oauth_access_filters();install_oauth_access_filters()
        for n in names:
            logger=logging.getLogger(n)
            assert len(logger.filters)==1 and isinstance(logger.filters[0],OAuthAccessFilter)
    finally:
        for n in names:logging.getLogger(n).filters=saved[n]
    for n in names:assert logging.getLogger(n).filters==saved[n]


@pytest.mark.parametrize('filename,disposition',[
 ('Реестр ЭЗС 2026.xlsx','attachment; filename="Reestr EZS 2026.xlsx"; filename*=UTF-8\'\'%D0%A0%D0%B5%D0%B5%D1%81%D1%82%D1%80%20%D0%AD%D0%97%D0%A1%202026.xlsx'),
 ('report.xlsx','attachment; filename="report.xlsx"; filename*=UTF-8\'\'report.xlsx'),
],ids=['unicode','ascii'])
def test_real_workbook_response_private_headers_cells_and_filename(filename,disposition):
    book=Workbook();sheet=book.active;sheet.title='Синтетика'
    sheet.append(['Показатель','Значение']);sheet.append(['Энергия, кВт·ч',12.5]);sheet.append(['Число',7])
    try:response=xlsx_response(book,filename)
    finally:book.close()
    assert response.status_code==200
    assert response.headers['cache-control']=='private, no-store'
    assert response.headers['content-type']==XLSX_MIME
    assert response.headers['content-disposition']==disposition
    assert response.body.startswith(b'PK\x03\x04') and 1000<len(response.body)<32768
    parsed=load_workbook(io.BytesIO(response.body),data_only=True,read_only=True)
    try:assert list(parsed['Синтетика'].values)==[('Показатель','Значение'),('Энергия, кВт·ч',12.5),('Число',7)]
    finally:parsed.close()


@pytest.mark.skipif(os.environ.get('RUN_PROTOCOL_PRIVACY_TESTS') != '1', reason='ROOT-only immutable boundary controls')
@pytest.mark.parametrize('control',['socket','dns','process','env','secret-file','outside-write','app-import','sql-import'],
                         ids=['socket','dns','process','env','secret-file','outside-write','app-import','sql-import'])
def test_offline_boundary_denial_controls(request,control):
    boundary=request.config._protocol_privacy_boundary
    boundary.denial_control(control)
