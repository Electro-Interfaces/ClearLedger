"""Регрессии сквозного аудита: ключи источника, повторная загрузка и время API."""
import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from sqlalchemy import select

from app.models import ChargeSession, Company, ServiceLocation
from app.services.asuim_normalize import _read_sheet, ingest_asuim_book, map_sessions
from app.services.charge_sessions_normalize import finalize_pending_session, ingest_charge_sessions
from app.services.network_reliability import network_reliability
from app.services.network_state import network_state
from app.services.ops_monitor_context import session_time_iso
from app.services.stations_normalize import build_source_station_index, resolve_session_location


def source_row(**changes):
    return {"id_сессии": "100", "id_станции": "3503", "номер_станции": "590",
            "телефон": "8 (900) 000-00-01", "id_пользователя": "42",
            "дата_начала": datetime(2026, 9, 17, 18),
            "дата_окончания": datetime(2026, 9, 18, 3, 43, 23),
            "статус": "Complete", "энергия_квтч": 12.9, "стоимость_руб": 179.31,
            "цена_тарифа": 13.9, **changes}


def test_source_phone_aliases_and_stable_station_identity():
    row = map_sessions([source_row()])[0]
    assert row['user_id'] == '+79000000001'
    assert row['source_station_id'] == '3503'
    assert row['station_code'] == '590'  # исторический номер сохраняется
    assert map_sessions([source_row(телефон=None, телефон_пользователя='+79000000002')])[0]['user_id'] == '+79000000002'
    assert map_sessions([source_row(телефон=None)])[0]['user_id'] is None
    assert map_sessions([source_row(телефон=None)])[0]['source_customer_id'] == '42'
    short = source_row(дата_окончания=datetime(2026, 9, 17, 17, 59, 59))
    assert map_sessions([short])[0]['duration_min'] == 0
    assert resolve_session_location(row, {'590': 'wrong'}, {'3503': 'correct'}) == 'correct'
    assert resolve_session_location(row, {'590': 'wrong'}, {'3503': None}) is None
    assert resolve_session_location(row, {'590': 'fallback'}, {}) == 'fallback'


def test_old_pending_rows_survive_incremental_day_filter():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['id_сессии', 'дата_начала'])
    ws.append([100, datetime(2026, 9, 17)])
    ws.append([101, datetime(2026, 9, 17)])
    ws.append([102, datetime(2026, 10, 21)])
    assert [r['id_сессии'] for r in _read_sheet(ws, 'дата_начала', '2026-10-20', {'100'})] == [100, 102]
    assert [r['id_сессии'] for r in _read_sheet(ws, 'дата_начала', '2026-10-20')] == [102]


@pytest.mark.parametrize('change', [
    {'finished_at': None}, {'result': 'Charging'},
    {'started_at': datetime(2026, 9, 16)},
])
def test_incomplete_or_different_source_cannot_finalize(change):
    row = {**map_sessions([source_row()])[0], **change}
    session = SimpleNamespace(finished_at=None, result='Charging', started_at=datetime(2026, 9, 17, 18))
    assert finalize_pending_session(session, row) is False
    assert session.result == 'Charging'


def test_finalization_keeps_accounting_enrichment_and_cannot_revert():
    row = map_sessions([source_row()])[0]
    session = SimpleNamespace(finished_at=None, result='Charging', started_at=row['started_at'],
                              client_tariff=Decimal('10'), client_name='Организация', user_id='original')
    assert finalize_pending_session(session, row)
    assert session.client_name == 'Организация' and session.user_id == 'original'
    assert session.client_amount == Decimal('129.0')
    assert session.amount == 179.31
    assert finalize_pending_session(session, row) is False


def test_session_time_has_explicit_moscow_offset():
    assert session_time_iso(None) is None
    assert session_time_iso(datetime(2026, 10, 9, 22, 50, 2)) == '2026-10-09T22:50:02+03:00'
    assert session_time_iso(datetime(2026, 10, 9, 19, 50, 2, tzinfo=timezone.utc)) == '2026-10-09T22:50:02+03:00'


async def fleet(db):
    company = Company(name='Проверка целостности', slug='integrity-' + uuid.uuid4().hex, profile_id='energy')
    foreign = Company(name='Чужой парк', slug='integrity-' + uuid.uuid4().hex, profile_id='energy')
    db.add_all([company, foreign])
    await db.flush()
    locations = []
    for cid, number, sid in [(company.id, '251-0075', '3503'),
                              (company.id, '590', 'else'), (foreign.id, 'foreign', '3503')]:
        loc = ServiceLocation(id='mi-' + uuid.uuid4().hex, company_id=cid,
                              code=uuid.uuid4().hex, station_number=number, name=number,
                              type='ev_charging', operational_status='working',
                              extra_metadata={'asuimStationId': sid, 'number': number})
        db.add(loc)
        locations.append(loc)
    await db.flush()
    return company.id, locations


@pytest.mark.asyncio(loop_scope='session')
async def test_import_resolves_stable_id_finalizes_once_and_preserves_final_history(db):
    cid, (correct, wrong, foreign) = await fleet(db)
    row = map_sessions([source_row()])[0]
    pending = {**row, 'finished_at': None, 'result': 'Charging', 'duration_min': 0, 'energy_kwh': 9.2, 'amount': 127.88}
    first = await ingest_charge_sessions(db, cid, [pending])
    assert first['created'] == 1 and first['errors'] == 0
    stored = (await db.execute(select(ChargeSession).where(ChargeSession.company_id == cid))).scalar_one()
    assert stored.location_id == correct.id and stored.station_code == '590'
    assert stored.user_id == '+79000000001'
    stored.client_name = 'Сохранённое обогащение'
    second = await ingest_charge_sessions(db, cid, [row])
    assert second['created'] == 0 and second['updated'] == 1 and second['errors'] == 0
    assert stored.result == 'Complete' and stored.client_name == 'Сохранённое обогащение'
    assert Decimal(str(stored.energy_kwh)) == Decimal('12.9')
    third = await ingest_charge_sessions(db, cid, [{**row, 'energy_kwh': 100}])
    assert third['updated'] == 0 and third['skipped'] == 1
    assert Decimal(str(stored.energy_kwh)) == Decimal('12.9')
    index = await build_source_station_index(db, cid)
    assert index['3503'] == correct.id and foreign.id not in index.values()
    wrong.extra_metadata = {'asuimStationId': '3503'}
    await db.flush()
    assert (await build_source_station_index(db, cid))['3503'] is None


@pytest.mark.asyncio(loop_scope='session')
async def test_book_includes_pending_outside_incremental_tail(db):
    cid, (loc, _, _) = await fleet(db)
    row = map_sessions([source_row()])[0]
    db.add_all([
        ChargeSession(company_id=cid, session_ext_id='100', location_id=loc.id,
                      started_at=row['started_at'], result='Charging'),
        ChargeSession(company_id=cid, session_ext_id='recent', location_id=loc.id,
                      started_at=datetime(2026, 10, 21), result='Complete'),
    ])
    await db.flush()
    wb = openpyxl.Workbook()
    ws = wb.active
    source = source_row()
    ws.append(list(source))
    ws.append(list(source.values()))
    payments = wb.create_sheet('payments')
    payments.append(['id_платежа', 'id_сессии', 'сумма_холда_руб.', 'дата'])
    content = io.BytesIO()
    wb.save(content)
    captured = []
    async def ingest(db, company, view, rows, **kw):
        if view == 'sessions':
            captured.extend(rows)
        return {'status': 'success'}
    with patch('app.services.asuim_normalize._ingest_view', ingest):
        await ingest_asuim_book(db, cid, content.getvalue())
    assert [str(r['id_сессии']) for r in captured] == ['100']


@pytest.mark.asyncio(loop_scope='session')
async def test_dry_attempts_only_after_last_success_with_history_and_company_scope(db):
    cid, (a, b, foreign) = await fleet(db)
    # До успеха — две неудачи, после — одна; будущее не влияет на выбранный день.
    for loc, events in [(a, [(8, 0), (9, 0), (10, 5), (11, 0), (13, 0)]),
                        (b, [(8, 0), (9, 0)]), (foreign, [(12, 100)])]:
        for day, energy in events:
            db.add(ChargeSession(company_id=loc.company_id, session_ext_id=uuid.uuid4().hex,
                                 location_id=loc.id, station_code=loc.station_number,
                                 started_at=datetime(2026, 10, day, 22, 50, 2),
                                 energy_kwh=energy, amount=0, result='Complete'))
    await db.flush()
    state = await network_state(db, cid, as_of=date(2026, 10, 12))
    rows = {r['locationId']: r for r in state['stations']}
    assert rows[a.id]['dryAttempts'] == 1
    assert rows[a.id]['lastDryAt'] == '2026-10-11T22:50:02+03:00'
    assert rows[a.id]['lastSessionAt'] == '2026-10-10T22:50:02+03:00'
    assert rows[b.id]['dryAttempts'] == 2  # станция без единой успешной зарядки
    assert foreign.id not in rows
    state = await network_state(db, cid, as_of=date(2026, 10, 10))
    assert next(r for r in state['stations'] if r['locationId'] == a.id)['dryAttempts'] == 0
    reliability = await network_reliability(db, cid, min_sessions=1, as_of=date(2026, 10, 12))
    assert next(r for r in reliability['stations'] if r['locationId'] == a.id)['lastSessionAt'].endswith('+03:00')


@pytest.mark.asyncio(loop_scope='session')
async def test_corporate_auto_enrichment_is_missing_only_and_matches_plan(db):
    from app.models import CorporateClient
    from app.services.charge_sessions_normalize import missing_org_enrichment
    cid, (loc, _, _) = await fleet(db)
    configs = [
        {'phone': '+79000000101', 'name': 'Тариф фиксированный', 'mode': 'flat', 'rate': Decimal('10')},
        {'phone': '+79000000102', 'name': 'Розница со скидкой', 'mode': 'retail', 'discount': Decimal('25')},
        {'phone': '+79000000103', 'name': 'Матрица', 'mode': 'matrix', 'discount': Decimal('20')},
    ]
    matrix = {'Матрица': {'Москва': {'CCS2': 20}}}
    for org in configs:
        db.add(CorporateClient(company_id=cid, phone=org['phone'], name=org['name'], mode=org['mode'],
                               rate=org.get('rate'), discount_pct=org.get('discount', 0),
                               matrix=matrix.get(org['name'])))
    rows = []
    plans = []
    for index, org in enumerate(configs):
        row = map_sessions([source_row(id_сессии=str(index), телефон=org['phone'],
                           тип_пользователя='ЮЛ', стоимость_руб=0, цена_тарифа=13.9,
                           энергия_квтч=12.345, регион='Москва', тип_коннектора='CCS Combo 2')])[0]
        rows.append(row)
        plans.append(missing_org_enrichment(SimpleNamespace(
            **{**row, 'client_name': None, 'client_tariff': None, 'client_amount': None}), org, matrix))
    result = await ingest_charge_sessions(db, cid, rows)
    assert result['corporate_priced'] == 3
    saved = list((await db.execute(select(ChargeSession).where(
        ChargeSession.company_id == cid).order_by(ChargeSession.session_ext_id))).scalars())
    for session, plan in zip(saved, plans):
        await db.refresh(session)
        for field in ('client_name', 'client_tariff', 'client_amount'):
            assert getattr(session, field) == plan[field]
    saved[0].client_name = 'Ручное обогащение'
    saved[0].client_tariff = Decimal('99')
    saved[0].client_amount = Decimal('100')
    repeated = await ingest_charge_sessions(db, cid, rows)
    assert repeated['corporate_priced'] == 0
    assert saved[0].client_amount == Decimal('100')
    assert saved[0].client_name == 'Ручное обогащение'


@pytest.mark.asyncio(loop_scope='session')
async def test_failed_rebuild_cannot_commit_partial_import_or_corporate_prices(db):
    from app.models import CorporateClient
    cid, (loc, _, _) = await fleet(db)
    db.add(CorporateClient(company_id=cid, phone='+79000000001', name='Тестовая организация',
                           mode='flat', rate=10))
    await db.commit()
    row = map_sessions([source_row(тип_пользователя='ЮЛ', стоимость_руб=0)])[0]
    with patch('app.services.charge_visits.recompute_visits', AsyncMock(side_effect=RuntimeError('Сбой пересчёта'))):
        with pytest.raises(RuntimeError, match='Сбой пересчёта'):
            await ingest_charge_sessions(db, cid, [row])
    await db.rollback()
    assert not list((await db.execute(select(ChargeSession).where(ChargeSession.company_id == cid))).scalars())


@pytest.mark.asyncio(loop_scope='session')
@pytest.mark.parametrize('shadow', ['test', 'closed', 'real'])
async def test_confirmed_old_number_beats_shadow_but_preserves_real_current_owner(db, shadow):
    from app.services.stations_normalize import build_station_index, backfill_session_locations
    cid, (actual, duplicate, _) = await fleet(db)
    actual.extra_metadata = {'asuimStationId': '3503', 'number': '251-0075',
                             'numberHistory': [{'было': '590'}]}
    if shadow == 'test':
        duplicate.is_test = True
        duplicate.extra_metadata = {'number': '590'}
    elif shadow == 'closed':
        duplicate.status = 'closed'
        duplicate.extra_metadata = {'number': '590'}
    await db.flush()
    expected = duplicate.id if shadow == 'real' else actual.id
    assert (await build_station_index(db, cid))['590'] == expected
    session = ChargeSession(company_id=cid, session_ext_id='legacy', station_code='590',
                            location_id=duplicate.id, started_at=datetime(2026, 10, 1))
    db.add(session)
    await db.flush()
    await backfill_session_locations(db, cid)
    await db.refresh(session)
    assert session.location_id == expected


@pytest.mark.asyncio(loop_scope='session')
async def test_historical_number_comes_from_source_id_not_wrong_session_location(db):
    from app.services.stations_normalize import build_station_index, backfill_session_locations
    cid, (actual, duplicate, _) = await fleet(db)
    duplicate.is_test = True
    duplicate.extra_metadata = {'number': '590'}
    # У реального объекта уже новый номер, истории в старом JSON не осталось.
    actual.extra_metadata = {'asuimStationId': '3503'}
    known = ChargeSession(company_id=cid, session_ext_id='proven', station_code='590',
                          source_station_id='3503', location_id=duplicate.id,
                          started_at=datetime(2026, 10, 1))
    legacy = ChargeSession(company_id=cid, session_ext_id='legacy', station_code='590',
                           location_id=duplicate.id, started_at=datetime(2026, 10, 2))
    db.add_all([known, legacy])
    await db.flush()
    assert (await build_station_index(db, cid))['590'] == actual.id
    await backfill_session_locations(db, cid)
    await db.refresh(known)
    await db.refresh(legacy)
    assert known.location_id == legacy.location_id == actual.id


@pytest.mark.asyncio(loop_scope='session')
async def test_number_alone_does_not_relocate_existing_real_object(db):
    from app.services.stations_normalize import build_station_index, backfill_session_locations
    cid, (actual, previous, _) = await fleet(db)
    current = ChargeSession(company_id=cid, session_ext_id='current', station_code='old-number',
                            source_station_id='3503', location_id=actual.id,
                            started_at=datetime(2026, 10, 2))
    historical = ChargeSession(company_id=cid, session_ext_id='history', station_code='old-number',
                               location_id=previous.id, started_at=datetime(2026, 1, 1))
    db.add_all([current, historical])
    await db.flush()
    assert (await build_station_index(db, cid))['old-number'] == actual.id
    await backfill_session_locations(db, cid)
    await db.refresh(historical)
    assert historical.location_id == previous.id


@pytest.mark.asyncio(loop_scope='session')
async def test_missing_phone_uses_company_registry_and_never_external_id_as_phone(db):
    from app.models import EzsCustomer
    cid, (loc, _, foreign) = await fleet(db)
    db.add_all([
        EzsCustomer(company_id=cid, customer_ext_id='42', phone='+79000000001'),
        EzsCustomer(company_id=foreign.company_id, customer_ext_id='42', phone='+79000000099'),
    ])
    await db.flush()
    rows = map_sessions([
        source_row(id_сессии='registry', телефон=None),
        source_row(id_сессии='explicit'),
        source_row(id_сессии='unknown', телефон=None, id_пользователя='999'),
    ])
    result = await ingest_charge_sessions(db, cid, rows)
    assert result['created'] == 3 and result['errors'] == 0
    stored = {s.session_ext_id: s for s in (await db.execute(select(ChargeSession).where(
        ChargeSession.company_id == cid))).scalars()}
    assert stored['registry'].user_id == stored['explicit'].user_id == '+79000000001'
    assert stored['unknown'].user_id is None
    assert stored['unknown'].source_customer_id == '999'
