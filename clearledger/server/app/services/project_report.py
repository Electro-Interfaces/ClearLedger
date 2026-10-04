"""Отчёты по проектам: история, чек-лист, работа — в Excel реестра и в отчёт проекта.

Один сборщик на два формата. Excel проекта и PDF-презентация берут данные отсюда же,
иначе цифры в файле и на слайде разъезжаются при первой же правке. История в
выгрузке — та же, что во вкладке «История» карточки: событие и его изменения
«поле — было — стало», по строке на изменение (так её фильтруют в Excel).
"""
from __future__ import annotations

import io
import json
from datetime import date, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select

from app.models import EzsSite, EzsSiteEvent, User

EVENT_KINDS = {
    "stage": "Смена стадии", "edit": "Правка", "touch": "Касание", "note": "Заметка",
    "doc": "Документ", "gate": "Чек-лист", "forecast": "Прогноз", "import": "Импорт",
    "result": "Результат", "decision": "Решение",
}
MSK = timezone(timedelta(hours=3))


def _when(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, datetime):
        return (dt.astimezone(MSK) if dt.tzinfo else dt).strftime("%d.%m.%Y %H:%M")
    return str(dt)


def _show(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)
    return text if len(text) <= 500 else text[:497] + "…"


def event_lines(e: EzsSiteEvent, stage_labels: dict[str, str], kind: str | None = None) -> list[dict[str, str]]:
    """Событие журнала — строками «поле — было — стало»; без изменений — одной строкой."""
    from app.services.ezs_site_work import relabel_stage_text
    text = relabel_stage_text(e.text, e.from_stage, e.to_stage, kind) if e.kind == "stage" else e.text
    base = {"kind": EVENT_KINDS.get(e.kind, e.kind), "text": text or ""}
    if e.kind == "stage" and (e.from_stage or e.to_stage) and not e.changes:
        return [{**base, "field": "Стадия", "old": stage_labels.get(e.from_stage or "", e.from_stage or ""),
                 "new": stage_labels.get(e.to_stage or "", e.to_stage or "")}]
    lines = []
    for c in e.changes or []:
        if not isinstance(c, dict):
            continue
        # Подпись стадии в журнале — снимок на момент записи, и у интеграции там
        # строительная («Оформление земли» вместо «Пилотное соглашение»): берём код.
        if c.get("field") == "stage" and (c.get("old") in stage_labels or c.get("new") in stage_labels):
            lines.append({**base, "field": "Стадия", "old": stage_labels.get(c.get("old") or "", c.get("old") or ""),
                          "new": stage_labels.get(c.get("new") or "", c.get("new") or "")})
            continue
        lines.append({**base, "field": c.get("label") or c.get("field") or "",
                      "old": _show(c.get("oldDisplay") if c.get("oldDisplay") is not None else c.get("old")),
                      "new": _show(c.get("newDisplay") if c.get("newDisplay") is not None else c.get("new"))})
    return lines or [{**base, "field": "", "old": "", "new": ""}]


async def history(db, company_id, site_ids=None, since: date | None = None, limit: int = 20000) -> list[dict[str, Any]]:
    from app.services.ezs_sites import STAGE_LABELS
    from app.services.ezs_checklist_integration import STAGE_LABELS as INT_LABELS
    labels_for = {"integration": {**STAGE_LABELS, **INT_LABELS}}
    q = (select(EzsSiteEvent, EzsSite.project_no, EzsSite.title, EzsSite.kind, User.name, User.email)
         .join(EzsSite, EzsSite.id == EzsSiteEvent.site_id)
         .outerjoin(User, User.id == EzsSiteEvent.author_user_id)
         .where(EzsSiteEvent.company_id == company_id))
    if site_ids is not None:
        q = q.where(EzsSiteEvent.site_id.in_(list(site_ids)))
    if since:
        q = q.where(EzsSiteEvent.created_at >= datetime.combine(since, datetime.min.time(), tzinfo=MSK))
    rows = (await db.execute(q.order_by(EzsSiteEvent.created_at.desc()).limit(limit))).all()
    out = []
    for e, no, title, kind, name, email in rows:
        labels = labels_for.get(kind or "", STAGE_LABELS)
        for line in event_lines(e, labels, kind):
            out.append({"at": _when(e.created_at), "projectNo": no or "", "title": title or "",
                        "author": name or email or ("система" if e.source == "system" else ""), **line})
    return out


def gate_line(site_no: str, title: str, stage_label: str, it: dict[str, Any]) -> dict[str, Any]:
    if it.get("needsConfirmation"):
        status = "Требует повторного подтверждения"
    elif it.get("done"):
        status = "Выполнено"
    elif it.get("waived"):
        status = "Обязательность снята"
    else:
        status = "Не выполнено"
    return {
        "projectNo": site_no, "title": title, "stage": stage_label, "key": it.get("key", ""),
        "label": it.get("label", ""), "role": it.get("role", ""),
        "required": "да" if it.get("required") else "", "status": status,
        "holds": "держит переход" if it.get("required") and not it.get("done") and not it.get("waived") else "",
        "by": it.get("confirmedBy") or it.get("waivedBy") or "",
        "at": _when(it.get("confirmedAt") or it.get("waivedAt") or ""),
        "note": it.get("waiveReason") or "",
    }


async def open_gates(db, sites: list[EzsSite]) -> list[dict[str, Any]]:
    """Реестр: незакрытые пункты текущей стадии — то, что держит проекты."""
    from app.services.ezs_site_work import gate_now
    out = []
    for s in sites:
        if s.stage in ("archive", "on_hold", "live"):
            continue
        g = await gate_now(db, s)
        for it in g.get("items") or []:
            if not it.get("done"):
                out.append(gate_line(s.project_no or "", s.title or "", g.get("stageLabel") or s.stage, it))
    return out


async def work_rows(db, company_id, user, site_ids) -> list[dict[str, Any]]:
    """Поручения и документы «Трека», привязанные к проектам (предмет `site:<id>`)."""
    from app.services import work_state
    from app.services.work_links import visible_work
    refs = {f"site:{i}": i for i in site_ids}
    if not refs:
        return []
    work = await visible_work(db, company_id, user)
    rows = (await db.execute(select(work).where(work.c.subject_ref.in_(list(refs))))).mappings().all()
    sites = {str(i): (no, t) for i, no, t in (await db.execute(
        select(EzsSite.id, EzsSite.project_no, EzsSite.title).where(EzsSite.id.in_(list(refs.values()))))).all()}
    people = dict((await db.execute(select(User.id, User.name).where(
        User.id.in_([r["responsible_id"] for r in rows if r["responsible_id"]] or [None])))).all())
    now = datetime.now(timezone.utc)
    out = []
    for r in rows:
        no, title = sites.get(r["subject_ref"].split(":", 1)[1], ("", ""))
        due = r["due_at"]
        out.append({
            "projectNo": no or "", "title": title or "", "kind": "Документ" if r["kind"] == "doc" else "Поручение",
            "key": r["key"] or "", "work": r["title"] or "", "state": work_state.COLUMN_NAMES.get(r["state"], r["state"]),
            "responsible": people.get(r["responsible_id"]) or "", "due": _when(due),
            "overdue": "просрочено" if due and due < now and r["state"] != "done" else "", "created": _when(r["at"]),
        })
    out.sort(key=lambda x: (x["projectNo"], x["state"] == "Готово", x["due"]))
    return out


# ── Отчёт одного проекта ────────────────────────────────────────────────────

async def project_report(db, company_id, user, site: EzsSite) -> dict[str, Any]:
    from app.services import ezs_sites, project_work
    from app.services.ezs_site_work import gate_state, site_doc_kinds, site_equipment_supplied, GATES_BY_KIND, GATES
    from app.services.ezs_project import TC_LABELS, EQ_LABELS, COST_LABELS, COST_CAPITAL, DOC_LABELS
    from sqlalchemy import text as sql

    detail = await ezs_sites.site_detail(db, company_id, site.id) or {}
    doc_kinds = await site_doc_kinds(db, site.id)
    equipment = await site_equipment_supplied(db, site.id)
    stages = [st for st in ezs_sites.STAGE_ORDER if GATES_BY_KIND.get(site.kind or "", GATES).get(st)]
    checklist, current = [], None
    for st in stages:
        g = gate_state(site, stage=st, doc_kinds=doc_kinds, equipment_supplied=equipment)
        if st == site.stage:
            current = g
        checklist += [gate_line(site.project_no or "", site.title or "", g.get("stageLabel") or st, it) for it in g.get("items") or []]
    work = await project_work.listing(db, company_id, user, site=site, scope="all", limit=500)
    work_items = [{
        "kind": "Документ" if w.get("kind") == "doc" else "Поручение", "key": w.get("key") or "",
        "work": w.get("title") or "", "state": w.get("state_name") or "", "responsible": w.get("responsible_name") or "",
        "due": _when(w.get("due_at")), "overdue": "просрочено" if w.get("overdue") else "", "created": _when(w.get("at")),
    } for w in work["items"]]
    p = {"sid": str(site.id), "cid": str(company_id)}
    tc = (await db.execute(sql("""select status, grid_operator, application_no, application_date, specs_no, specs_date,
        contract_no, contract_date, power_kwt, cost, total_cost, due_date, done_date from ezs_tech_connections
        where site_id = :sid and company_id = :cid"""), p)).mappings().first()
    eq = (await db.execute(sql("""select title, manufacturer, power_kwt, qty, status, supplier, price, due_date,
        supplied_date, installed_date, serial_number from ezs_site_equipment where site_id = :sid and company_id = :cid
        and status <> 'cancelled' order by created_at"""), p)).mappings().all()
    costs = (await db.execute(sql("""select kind, title, plan_amount, fact_amount, doc_ref from ezs_site_costs
        where site_id = :sid and company_id = :cid order by created_at"""), p)).mappings().all()
    docs = (await db.execute(sql("""select d.kind, d.title, d.stage, d.created_at, u.name from ezs_site_docs d
        left join users u on u.id = d.uploaded_by where d.site_id = :sid and d.company_id = :cid
        order by d.created_at desc"""), p)).mappings().all()
    hist = await history(db, company_id, [site.id])
    verdict = None
    if site.kind != "integration" and site.stage in ezs_sites.STAGE_ORDER:
        from app.services.ezs_site_analysis import QUADRANTS, verdicts
        v = (await verdicts(db, company_id, [site])).get(str(site.id))
        if v:
            q = QUADRANTS.get(v["quadrant"], {})
            verdict = {"label": q.get("label", v["quadrant"]), "hint": q.get("hint", ""),
                       "confidence": v.get("confidence"), "unknown": v.get("unknown") or []}
    plan = sum(float(c["plan_amount"] or 0) for c in costs)
    fact = sum(float(c["fact_amount"] or 0) for c in costs)
    report: dict[str, Any] = {
        "generatedAt": _when(datetime.now(timezone.utc)),
        "summary": {
            "projectNo": site.project_no, "title": site.title, "kind": site.kind or "new_build",
            "stage": site.stage, "stageLabel": detail.get("stageLabel") or site.stage,
            "phaseLabel": detail.get("phaseLabel") or "", "owner": detail.get("ownerName") or "",
            "region": site.region_norm or site.region or "", "address": site.full_address or site.address or "",
            "nextAction": site.next_action or "", "nextActionDue": site.next_action_due or "",
            "stageSince": str(site.stage_since or ""), "createdAt": _when(site.created_at),
            "plannedPowerKwt": float(site.planned_power_kwt) if site.planned_power_kwt is not None else None,
            "plannedEzs": site.planned_ezs_count, "freePowerKwt": site.free_power_num,
            "controlForm": site.control_form or "", "contractStart": site.contract_start or "",
            "commissionedOn": site.commissioned_on or "", "exitKind": site.exit_kind or "",
        },
        "verdict": verdict,
        "gate": {"stageLabel": (current or {}).get("stageLabel") or "", "blocking": (current or {}).get("blocking") or [],
                 "closed": sum(1 for i in (current or {}).get("items") or [] if i.get("done")),
                 "total": len((current or {}).get("items") or [])},
        "phases": phase_strip(site),
        "checklist": checklist, "work": work_items,
        "techConnection": None if not tc else {
            "status": TC_LABELS.get(tc["status"] or "", tc["status"] or ""), "gridOperator": tc["grid_operator"] or "",
            "applicationNo": tc["application_no"] or "", "applicationDate": str(tc["application_date"] or ""),
            "specsNo": tc["specs_no"] or "", "specsDate": str(tc["specs_date"] or ""),
            "contractNo": tc["contract_no"] or "", "contractDate": str(tc["contract_date"] or ""),
            "powerKwt": float(tc["power_kwt"] or 0) or None, "cost": float(tc["cost"] or 0) or None,
            "totalCost": float(tc["total_cost"] or 0) or None, "due": str(tc["due_date"] or ""), "done": str(tc["done_date"] or "")},
        "equipment": [{"title": r["title"] or "", "manufacturer": r["manufacturer"] or "", "powerKwt": float(r["power_kwt"] or 0) or None,
                       "qty": r["qty"], "status": EQ_LABELS.get(r["status"] or "", r["status"] or ""), "supplier": r["supplier"] or "",
                       "price": float(r["price"] or 0) or None, "due": str(r["due_date"] or ""), "supplied": str(r["supplied_date"] or ""),
                       "installed": str(r["installed_date"] or ""), "serial": r["serial_number"] or ""} for r in eq],
        "costs": [{"kind": COST_LABELS.get(r["kind"], r["kind"]), "capital": "капвложение" if COST_CAPITAL.get(r["kind"]) else "расход периода",
                   "title": r["title"] or "", "plan": float(r["plan_amount"] or 0) or None, "fact": float(r["fact_amount"] or 0) or None,
                   "docRef": r["doc_ref"] or ""} for r in costs],
        "budget": {"plan": round(plan, 2), "fact": round(fact, 2), "delta": round(fact - plan, 2)},
        "documents": [{"kind": DOC_LABELS.get(r["kind"], r["kind"] or ""), "title": r["title"] or "", "stage": r["stage"] or "",
                       "at": _when(r["created_at"]), "by": r["name"] or ""} for r in docs],
        "history": hist,
    }
    if site.kind == "integration":
        report["integration"] = await integration_part(db, company_id, site)
    return report


def phase_strip(site) -> list[dict[str, str]]:
    """Этапы проекта для полосы на титуле: пройден, текущий, впереди."""
    from app.services.ezs_sites import phases_for, stage_phase
    phases = [p for p in phases_for(site.kind) if p["key"] != "closed"]
    current = stage_phase(site.stage, site.kind)
    keys = [p["key"] for p in phases]
    pos = keys.index(current) if current in keys else -1
    return [{"label": p["label"], "state": "done" if i < pos else "current" if i == pos else "next"} for i, p in enumerate(phases)]


async def integration_part(db, company_id, site) -> dict[str, Any]:
    from app.services import project_integration as I
    data = I.read(site)
    names = {s["id"]: s for s in await I.station_catalog(db, company_id)}
    labels = {"outgoing": "Наши ЭЗС у партнёра", "incoming": "ЭЗС партнёра у нас",
              "information": "информационная", "roaming": "роуминг"}
    stations = []
    for sc in data["scenarios"]:
        scen = sc.get("name") or f'{labels[sc["direction"]]} · {labels[sc["format"]]}'
        meta = sc.get("connectedMeta") or {}
        for sid in sc["selectedIds"]:
            st = names.get(sid) or {}
            m = meta.get(sid) or {}
            stations.append({"scenario": scen, "code": st.get("code") or sid, "name": st.get("name") or "нет в справочнике",
                             "region": st.get("region") or "", "city": st.get("city") or "",
                             "agreed": "да" if sid in sc["agreedIds"] else "", "pilot": "да" if sid in sc["pilotIds"] else "",
                             "connected": "да" if sid in sc["connectedIds"] else "", "connectedAt": m.get("at", ""),
                             "basis": {"check": "проверено у принимающей стороны", "session": "первая сессия"}.get(m.get("basis", ""), ""),
                             "retired": "выбыла" if st.get("lifecycle") == "closed" or st.get("opStatus") == "decommissioned" else ""})
    payer = {"partner": "партнёр платит нам", "us": "мы платим партнёру", "none": "без расчётов", "": ""}
    return {
        "partner": data["partner"], "commercial": data["commercial"], "settlement": data["settlement"],
        "technical": data["technical"], "work": data["work"],
        "scenarios": [{"name": sc.get("name") or f'{labels[sc["direction"]]} · {labels[sc["format"]]}',
                       "direction": labels[sc["direction"]], "format": labels[sc["format"]],
                       "payer": payer.get(sc.get("payer") or "", ""), "model": sc.get("model") or "", "rate": sc.get("rate") or "",
                       "clientPrice": sc.get("clientPrice") or "",
                       "selected": len(sc["selectedIds"]), "agreed": len(sc["agreedIds"]), "connected": len(sc["connectedIds"]),
                       "pilot": len(sc["pilotIds"])} for sc in data["scenarios"]],
        "stations": stations,
        "tests": [{"title": t["title"], "required": "да" if t["required"] else "",
                   "status": {"pending": "не проведено", "passed": "пройдено", "failed": "замечание", "na": "неприменимо"}[t["status"]],
                   "session": t.get("sessionRef") or "", "comment": t.get("comment") or "", "by": t.get("byName") or "", "at": _when(t.get("at"))}
                  for t in data["tests"]],
        "reconciliations": [{"kind": "пробная по пилоту" if r["kind"] == "pilot" else "месячная", "period": r["period"],
                             "ourSessions": r["ours"]["sessions"], "partnerSessions": r["partner"]["sessions"],
                             "ourKwh": r["ours"]["kwh"], "partnerKwh": r["partner"]["kwh"],
                             "ourAmount": r["ours"]["amount"], "partnerAmount": r["partner"]["amount"],
                             "state": {"match": "сходится", "resolved": "урегулировано", "diff": "расхождение"}[I.recon_state(r)],
                             "resolution": r.get("resolution") or "", "by": r.get("byName") or ""} for r in data["reconciliations"]],
        "listVersions": [{"version": v.get("version"), "stations": len(v.get("stationIds") or []), "note": v.get("note") or "",
                          "by": v.get("byName") or "", "at": _when(v.get("at"))} for v in data["listVersions"]],
    }


# ── Excel ───────────────────────────────────────────────────────────────────

def add_sheet(wb, title: str, columns: list[tuple[str, str, int]], rows: list[dict[str, Any]]):
    """Лист-таблица: заголовок жирным, закреплён, ширины колонок, автофильтр."""
    from openpyxl.styles import Alignment, Font
    ws = wb.create_sheet(title[:31])
    ws.append([c[1] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(c[0]) for c in columns])
    for i, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


TC_FIELDS = [("status", "Статус"), ("gridOperator", "Сетевая организация"), ("applicationNo", "Заявка №"),
             ("applicationDate", "Заявка от"), ("specsNo", "ТУ №"), ("specsDate", "ТУ от"), ("contractNo", "Договор ТП №"),
             ("contractDate", "Договор от"), ("powerKwt", "Мощность, кВт"), ("cost", "Договор с сетевой, ₽"),
             ("totalCost", "Итого по ТУ, ₽"), ("due", "Срок мероприятий"), ("done", "Исполнено")]
HISTORY_COLUMNS = [("at", "Когда", 17), ("projectNo", "Проект", 16), ("title", "Название", 28), ("author", "Кто", 22),
                   ("kind", "Событие", 14), ("text", "Описание", 40), ("field", "Что изменилось", 26),
                   ("old", "Было", 28), ("new", "Стало", 28)]
CHECKLIST_COLUMNS = [("projectNo", "Проект", 16), ("title", "Название", 28), ("stage", "Стадия", 18), ("key", "Пункт", 7),
                     ("label", "Требование", 60), ("role", "Кто отвечает", 14), ("required", "Обязательный", 12),
                     ("status", "Состояние", 20), ("holds", "Переход", 14), ("by", "Кто подтвердил / снял", 22),
                     ("at", "Когда", 17), ("note", "Обоснование", 30)]
WORK_COLUMNS = [("projectNo", "Проект", 16), ("title", "Название", 28), ("kind", "Вид", 12), ("key", "Номер", 10),
                ("work", "Работа", 50), ("state", "Состояние", 16), ("responsible", "Исполнитель", 22),
                ("due", "Срок", 17), ("overdue", "Просрочка", 12), ("created", "Создано", 17)]


def project_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    s = report["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    pairs = [
        ("Проект", f'{s["projectNo"] or ""} · {s["title"] or ""}'), ("Этап / стадия", f'{s["phaseLabel"]} · {s["stageLabel"]}'.strip(" ·")),
        ("В стадии с", s["stageSince"]), ("Ответственный", s["owner"] or "не назначен"),
        ("Следующий шаг", s["nextAction"]), ("Срок шага", s["nextActionDue"]),
        ("Решение", (lambda v: f'{v["label"]} — {v["hint"]}' if v else "")(report.get("verdict"))),
        ("Чек-лист текущей стадии", f'закрыто {report["gate"]["closed"]} из {report["gate"]["total"]}'),
        ("Держит переход", "; ".join(report["gate"]["blocking"])),
        ("Регион", s["region"]), ("Адрес", s["address"]),
    ]
    if s["kind"] != "integration":
        pairs += [("Мощность план, кВт", s["plannedPowerKwt"]), ("ЭЗС план", s["plannedEzs"]),
                  ("Свободная мощность", s["freePowerKwt"]), ("Форма оформления", s["controlForm"]),
                  ("Договор с", s["contractStart"]), ("Введён", s["commissionedOn"])]
    pairs += [("Бюджет план, ₽", report["budget"]["plan"] or None), ("Бюджет факт, ₽", report["budget"]["fact"] or None),
              ("Создан", s["createdAt"]), ("Отчёт сформирован", report["generatedAt"])]
    integ = report.get("integration")
    if integ:
        pairs += [("Партнёр", integ["partner"].get("name", "")), ("Юрлицо партнёра", integ["partner"].get("legalEntity", "")),
                  ("Цель", integ["partner"].get("purpose", ""))]
    for k, v in pairs:
        ws.append([k, v])
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    add_sheet(wb, "Чек-лист", CHECKLIST_COLUMNS[2:], report["checklist"])
    add_sheet(wb, "Работа", WORK_COLUMNS[2:], report["work"])
    if integ:
        add_sheet(wb, "Сценарии", [("name", "Сценарий", 30), ("direction", "Направление", 22), ("format", "Формат", 16),
                  ("payer", "Расчёты", 22), ("model", "Модель", 14), ("rate", "Ставка", 18), ("clientPrice", "Цена клиенту", 20),
                  ("selected", "Выбрано", 10), ("agreed", "Согласовано", 12), ("connected", "Подключено", 12), ("pilot", "Пилот", 8)],
                  integ["scenarios"])
        add_sheet(wb, "Станции", [("scenario", "Сценарий", 28), ("code", "Код", 10), ("name", "Станция", 32), ("region", "Регион", 22),
                  ("city", "Город", 16), ("agreed", "Согласована", 12), ("pilot", "Пилот", 8), ("connected", "Подключена", 12),
                  ("connectedAt", "Дата подключения", 16), ("basis", "Основание", 26), ("retired", "Выбыла", 10)], integ["stations"])
        add_sheet(wb, "Испытания", [("title", "Испытание", 50), ("required", "Обязательное", 12), ("status", "Статус", 14),
                  ("session", "Тестовая сессия", 18), ("comment", "Комментарий", 36), ("by", "Кто", 22), ("at", "Когда", 17)], integ["tests"])
        add_sheet(wb, "Сверки", [("kind", "Вид", 18), ("period", "Период", 18), ("ourSessions", "Сессии: мы", 11),
                  ("partnerSessions", "Сессии: партнёр", 14), ("ourKwh", "кВт·ч: мы", 11), ("partnerKwh", "кВт·ч: партнёр", 14),
                  ("ourAmount", "₽: мы", 12), ("partnerAmount", "₽: партнёр", 12), ("state", "Итог", 14),
                  ("resolution", "Урегулирование", 36), ("by", "Кто", 22)], integ["reconciliations"])
        add_sheet(wb, "Версии перечней", [("version", "Версия", 8), ("stations", "ЭЗС", 8), ("note", "Документ / комментарий", 40),
                  ("by", "Кто", 22), ("at", "Когда", 17)], integ["listVersions"])
        add_sheet(wb, "Паспорт интеграции", [("section", "Раздел", 26), ("field", "Поле", 30), ("value", "Значение", 80)],
                  [{"section": sec, "field": k, "value": v} for sec, block in (("Коммерческие условия", integ["commercial"]),
                   ("Порядок расчётов", integ["settlement"]), ("Технические параметры", integ["technical"]), ("Пилот", integ["work"]))
                   for k, v in block.items() if v])
    else:
        tc = report["techConnection"] or {}
        add_sheet(wb, "Присоединение", [("k", "Параметр", 28), ("v", "Значение", 40)],
                  [{"k": label, "v": tc.get(key)} for key, label in TC_FIELDS if tc.get(key) not in (None, "")])
        add_sheet(wb, "Оборудование", [("title", "Позиция", 30), ("manufacturer", "Производитель", 18), ("powerKwt", "кВт", 8),
                  ("qty", "Кол-во", 8), ("status", "Статус", 16), ("supplier", "Поставщик", 22), ("price", "Цена, ₽", 12),
                  ("due", "Срок", 12), ("supplied", "Поставлено", 12), ("installed", "Смонтировано", 12), ("serial", "Серийный №", 16)],
                  report["equipment"])
        add_sheet(wb, "Бюджет", [("kind", "Статья", 26), ("capital", "Судьба затрат", 16), ("title", "Описание", 32),
                  ("plan", "План, ₽", 14), ("fact", "Факт, ₽", 14), ("docRef", "Основание факта", 24)], report["costs"])
    add_sheet(wb, "Документы", [("kind", "Вид", 26), ("title", "Название", 40), ("stage", "Стадия", 14), ("at", "Приложен", 17),
              ("by", "Кто", 22)], report["documents"])
    add_sheet(wb, "История", HISTORY_COLUMNS[:1] + HISTORY_COLUMNS[3:], report["history"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
